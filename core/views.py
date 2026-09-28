import csv
import calendar
import hashlib
import hmac
import io
import json
import re
import secrets
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from urllib.parse import urlencode

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.tokens import default_token_generator
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.mail import send_mail, send_mass_mail
from django.db import IntegrityError, connection, transaction
from django.db.models import Count, Q
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode
from django.views.decorators.csrf import ensure_csrf_cookie
from openpyxl import Workbook, load_workbook

from .models import (
    AcademicPeriod, AcademicPlan, Attendance, Audit, Book, Claim, Course,
    Enrollment, Grade, GradingScale, ImportBatch, LostItem, Loan, Membership,
    Notice, NoticeRead, Offering, PlanSubject, PlanYear, PlatformBillingSettings,
    PlatformPriceChange, School, SchoolEvent, SchoolSignupRequest, SchoolSubscription, Section, SubscriptionPlan,
    Student, StudentGuardian, Subject, SubscriptionCharge, TeacherAssignment, User,
)


ROLE_LABELS = dict(Membership.Role.choices)
ROLE_ALIASES = {"directivo": "directivo", "secretaria": "secretaria", "preceptor": "preceptor", "docente": "docente", "biblioteca": "biblioteca", "alumno": "alumno", "tutor": "tutor", "school_admin": "school_admin"}
WRITE_ROLES = {"school_admin", "directivo"}
ATTENDANCE_ROLES = WRITE_ROLES | {"secretaria", "preceptor"}
GRADE_ROLES = WRITE_ROLES | {"docente"}
LIBRARY_ROLES = WRITE_ROLES | {"biblioteca"}
NOTICE_ROLES = WRITE_ROLES | {"secretaria", "preceptor"}
ACADEMIC_ROLES = WRITE_ROLES | {"docente"}


def json_body(request):
    try:
        return json.loads(request.body or b"{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        raise ValueError("El contenido no es un JSON válido.")


def error(message, status=400):
    if isinstance(message, IntegrityError):
        message = "No se pudo guardar. Revisá que no haya registros duplicados o datos incompatibles."
    elif isinstance(message, (KeyError, TypeError)):
        message = "Faltan datos o tienen un formato incorrecto."
    elif isinstance(message, Exception) and not isinstance(message, ValueError):
        message = "No se pudo completar la operación."
    return JsonResponse({"error": str(message)}, status=status)


def active_school(request):
    if not getattr(request, "active_school", None):
        raise PermissionError("La cuenta no tiene una escuela activa asignada.")
    return request.active_school


def active_role(request):
    if request.user.is_superuser:
        return "platform_admin"
    return request.school_membership.role if request.school_membership else None


def allows(request, *roles):
    return request.user.is_superuser or active_role(request) in roles


def school_required(view):
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return error("Iniciá sesión para continuar.", 401)
        if not request.active_school:
            return error("Seleccioná una escuela activa para continuar.", 409)
        if request.active_school.state == School.State.SUSPENDED and not request.user.is_superuser:
            return error("El acceso a esta escuela está suspendido.", 403)
        return view(request, *args, **kwargs)
    return wrapped


def api_methods(*methods):
    def decorator(view):
        @wraps(view)
        def wrapped(request, *args, **kwargs):
            if request.method not in methods:
                response = error("Método no permitido.", 405)
                response["Allow"] = ", ".join(methods)
                return response
            return view(request, *args, **kwargs)
        return wrapped
    return decorator


def audit(request, action, detail):
    if request.active_school:
        Audit.objects.create(school=request.active_school, actor=request.user if request.user.is_authenticated else None, action=action, detail=detail[:500])


def set_rls_school(school):
    """Set the trusted tenant scope for the current outer request transaction."""
    if connection.vendor == "postgresql":
        with connection.cursor() as cursor:
            cursor.execute("SELECT set_config('app.current_school_id', %s, true)", [str(school.pk)])


def school_membership(request, role=None):
    q = Membership.objects.filter(school=active_school(request), user=request.user, is_active=True)
    if role:
        q = q.filter(role=role)
    return q.first()


def membership_for(request, role):
    m = school_membership(request, role)
    if not m and not request.user.is_superuser:
        raise PermissionError("Tu perfil no tiene permiso para esta acción.")
    return m


def has_role(request, *roles):
    if not allows(request, *roles):
        raise PermissionError("Tu perfil no tiene permiso para esta acción.")


def can_see_student(request, student):
    role = active_role(request)
    if request.user.is_superuser or role in WRITE_ROLES | {"secretaria", "preceptor", "biblioteca"}:
        return True
    if role == "alumno":
        return student.account_id == request.user.id
    if role == "tutor":
        return StudentGuardian.objects.filter(school=student.school, student=student, guardian=request.user).exists()
    if role == "docente":
        return Enrollment.objects.filter(school=student.school, student=student, academic_year=timezone.localdate().year,
            state=Enrollment.State.ACTIVE, section__offerings__teacher_assignments__membership=request.school_membership).exists()
    return False


def student_query(request):
    school = active_school(request)
    qs = Student.objects.filter(school=school).prefetch_related("enrollments__section__plan_year")
    role = active_role(request)
    if role == "alumno":
        qs = qs.filter(account=request.user)
    elif role == "tutor":
        qs = qs.filter(guardian_links__guardian=request.user)
    elif role == "docente":
        qs = qs.filter(enrollments__school=school, enrollments__academic_year=timezone.localdate().year,
                       enrollments__state=Enrollment.State.ACTIVE,
                       enrollments__section__offerings__teacher_assignments__membership=request.school_membership).distinct()
    return qs.order_by("name")


def enrollment_for(student, school, year=None):
    year = year or timezone.localdate().year
    return Enrollment.objects.filter(school=school, student=student, academic_year=year, state=Enrollment.State.ACTIVE).select_related("section__plan_year").first()


def student_data(student, school):
    e = enrollment_for(student, school)
    return {"id": student.id, "name": student.name, "email": student.email,
            "source_id": student.source_id, "section_id": e.section_id if e else None,
            "course": e.section.plan_year.year_label if e else "",
            "division": e.section.division if e else "", "shift": e.section.get_shift_display() if e else "",
            "status": student.get_status_display()}


def audience_code(value):
    aliases = {"Todos": Notice.Audience.ALL, "Alumnos": Notice.Audience.STUDENTS,
               "Familias": Notice.Audience.GUARDIANS, "Docentes": Notice.Audience.TEACHERS,
               "Personal": Notice.Audience.STAFF}
    if value in aliases:
        return aliases[value]
    if value in Notice.Audience.values:
        return value
    return Notice.Audience.ALL


def audience_label(value):
    return dict(Notice.Audience.choices).get(value, value)


def notice_audiences(request):
    role = active_role(request)
    mapping = {
        "alumno": {Notice.Audience.ALL, Notice.Audience.STUDENTS},
        "tutor": {Notice.Audience.ALL, Notice.Audience.GUARDIANS},
        "docente": {Notice.Audience.ALL, Notice.Audience.TEACHERS, Notice.Audience.STAFF},
        "biblioteca": {Notice.Audience.ALL, Notice.Audience.STAFF},
        "preceptor": {Notice.Audience.ALL, Notice.Audience.STAFF},
        "secretaria": {Notice.Audience.ALL, Notice.Audience.STAFF},
        "directivo": {Notice.Audience.ALL, Notice.Audience.STAFF},
        "school_admin": set(Notice.Audience.values),
        "platform_admin": set(Notice.Audience.values),
    }
    return mapping.get(role, {Notice.Audience.ALL})


def serialize_notice(n, user=None):
    read_at = None
    if user and not user.is_superuser:
        receipt = next((item for item in n.reads.all() if item.user_id == user.id), None)
        read_at = receipt.read_at.isoformat() if receipt else None
    return {"id": n.id, "title": n.title, "body": n.body, "audience": audience_label(n.audience),
            "created_at": n.created_at.isoformat(), "read_at": read_at, "is_read": bool(read_at)}


def serialize_event(event):
    return {"id": event.id, "title": event.title, "description": event.description,
            "starts_at": timezone.localtime(event.starts_at).isoformat(),
            "ends_at": timezone.localtime(event.ends_at).isoformat() if event.ends_at else None,
            "audience": audience_label(event.audience)}


def serialize_attendance(a):
    return {"id": a.id, "student_id": a.student_id, "student": a.student.name,
            "course": enrollment_for(a.student, a.school).section.plan_year.year_label if enrollment_for(a.student, a.school) else "",
            "division": enrollment_for(a.student, a.school).section.division if enrollment_for(a.student, a.school) else "",
            "offering": a.offering.subject.name if a.offering_id else "Jornada",
            "date": a.date.isoformat(), "status": a.get_status_display(), "note": a.note}


def serialize_grade(g):
    return {"id": g.id, "student_id": g.student_id, "student": g.student.name,
            "subject": g.offering.subject.name, "period": g.period.name,
            "grade": float(g.value), "note": g.note}


def serialize_book(b):
    active = b.loans.filter(returned_at__isnull=True).count()
    return {"id": b.id, "title": b.title, "author": b.author, "code": b.code,
            "copies": b.copies, "available": max(0, b.copies-active)}


def serialize_loan(l):
    return {"id": l.id, "book_id": l.book_id, "student_id": l.student_id, "book": l.book.title,
            "student": l.student.name, "borrowed_at": timezone.localtime(l.borrowed_at).strftime("%Y-%m-%d %H:%M:%S"),
            "due_at": l.due_at.isoformat(), "returned_at": timezone.localtime(l.returned_at).strftime("%Y-%m-%d %H:%M:%S") if l.returned_at else None}


def serialize_lost(item):
    return {"id": item.id, "title": item.title, "description": item.description, "place": item.place,
            "status": item.get_status_display(), "created_at": item.created_at.isoformat()}


def serialize_claim(c):
    return {"id": c.id, "item_id": c.item_id, "item": c.item.title, "student_id": c.student_id,
            "student": c.student.name, "message": c.message, "created_at": c.created_at.isoformat(),
            "status": c.get_status_display()}


def serialize_plan(plan):
    return {"id": plan.id, "name": plan.name, "school_type": plan.school_type,
            "jurisdiction": plan.jurisdiction, "orientation": plan.orientation,
            "valid_from_year": plan.valid_from_year, "is_active": plan.is_active}


def serialize_year(y):
    return {"id": y.id, "plan_id": y.plan_id, "year_label": y.year_label,
            "ordinal": y.ordinal, "cycle": y.cycle, "orientation": y.orientation,
            "subjects": [{"id": s.subject_id, "name": s.subject.name,
                          "kind": s.subject.kind, "weekly_hours": float(s.weekly_hours)} for s in y.subject_entries.select_related("subject").all()]}


def serialize_section(s):
    return {"id": s.id, "plan_year_id": s.plan_year_id, "year_label": s.plan_year.year_label,
            "division": s.division, "shift": s.shift, "shift_label": s.get_shift_display(),
            "academic_year": s.academic_year, "name": s.name or f"{s.plan_year.year_label} {s.division}"}


def serialize_enrollment(enrollment):
    return {"id": enrollment.id, "student_id": enrollment.student_id, "student": enrollment.student.name,
            "section_id": enrollment.section_id, "section": serialize_section(enrollment.section),
            "academic_year": enrollment.academic_year, "state": enrollment.get_state_display(),
            "started_at": enrollment.started_at.isoformat(),
            "ended_at": enrollment.ended_at.isoformat() if enrollment.ended_at else None}


def serialize_offering(o):
    return {"id": o.id, "section_id": o.section_id, "section": serialize_section(o.section),
            "subject_id": o.subject_id, "subject": o.subject.name, "kind": o.subject.kind,
            "teachers": [x.membership.user.name for x in o.teacher_assignments.select_related("membership__user").all()]}


def page(request):
    return render(request, "core/index.html", {"demo_mode": settings.DEMO_MODE})


@ensure_csrf_cookie
def home(request):
    if request.user.is_authenticated and request.user.is_superuser:
        return redirect("platform_admin")
    return page(request)


@ensure_csrf_cookie
@login_required(login_url="/")
def platform_admin(request):
    if not request.user.is_superuser:
        return HttpResponseForbidden("Esta consola está disponible solo para administración de la plataforma.")
    return render(request, "core/platform_admin.html", {"name": request.user.name})


def api_login(request):
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        data = json_body(request)
    except ValueError as exc:
        return error(exc)
    user = authenticate(request, username=data.get("email", "").strip().lower(), password=data.get("password", ""))
    if user is None or not user.is_active:
        return error("Correo o contraseña incorrectos.", 401)
    login(request, user)
    membership = Membership.objects.filter(user=user, is_active=True, school__state__in=(School.State.TRIAL, School.State.ACTIVE)).select_related("school").order_by("school__name").first()
    if membership and not user.is_superuser:
        request.session["active_school_id"] = membership.school_id
    else:
        request.session.pop("active_school_id", None)
    return JsonResponse({"ok": True})


def api_logout(request):
    if request.method != "POST":
        return error("Método no permitido.", 405)
    logout(request)
    return JsonResponse({"ok": True})


def api_me(request):
    if not request.user.is_authenticated:
        return error("Iniciá sesión para continuar.", 401)
    memberships = Membership.objects.filter(user=request.user, is_active=True, school__state__in=(School.State.TRIAL, School.State.ACTIVE)).select_related("school").order_by("school__name")
    schools = [{"id": m.school_id, "name": m.school.name, "school_type": m.school.school_type,
                "role": m.role, "jurisdiction": m.school.jurisdiction, "theme": m.school.theme} for m in memberships]
    if request.user.is_superuser:
        schools = [{"id": s.id, "name": s.name, "school_type": s.school_type,
                    "role": "platform_admin", "jurisdiction": s.jurisdiction, "theme": s.theme}
                   for s in School.objects.exclude(state=School.State.SUSPENDED).order_by("name")]
    membership = request.school_membership
    role = "platform_admin" if request.user.is_superuser else membership.role if membership else None
    active = request.active_school
    subscription = getattr(active, "subscription", None) if active else None
    if subscription:
        apply_due_subscription_plan(subscription)
    current_plan = subscription.plan if subscription else SubscriptionPlan.BASIC
    analytics_enabled = bool(active and active.state == School.State.ACTIVE and subscription
        and subscription.state == SchoolSubscription.State.ACTIVE and current_plan == SubscriptionPlan.PRO)
    return JsonResponse({"id": request.user.id, "name": request.user.name, "email": request.user.email,
                         "role": role, "school_id": active.id if active else None,
                         "school_name": active.name if active else None,
                         "school_type": active.school_type if active else None,
                         "jurisdiction": active.jurisdiction if active else "",
                         "school_theme": active.theme if active else School.Theme.FOREST,
                         "subscription_plan": current_plan,
                         "subscription_state": subscription.state if subscription else None,
                         "analytics_enabled": analytics_enabled, "schools": schools})


def api_schools(request):
    if request.method == "GET":
        if not request.user.is_authenticated:
            return error("Iniciá sesión para continuar.", 401)
        if request.user.is_superuser:
            items = School.objects.all().order_by("name")
        else:
            items = School.objects.filter(memberships__user=request.user, memberships__is_active=True,
                                          state__in=(School.State.ACTIVE, School.State.TRIAL)).distinct().order_by("name")
        return JsonResponse([{"id": s.id, "name": s.name, "school_type": s.school_type, "state": s.state} for s in items], safe=False)
    if request.method == "POST":
        if not request.user.is_authenticated:
            return error("Iniciá sesión para continuar.", 401)
        try:
            data = json_body(request)
            sid = int(data.get("school_id"))
        except (ValueError, TypeError) as exc:
            return error(exc or "Escuela inválida.")
        school = School.objects.filter(pk=sid, state__in=(School.State.TRIAL, School.State.ACTIVE)).first()
        if not school:
            return error("La escuela no existe o está suspendida.", 404)
        membership = Membership.objects.filter(user=request.user, school=school, is_active=True).first()
        if not membership and not request.user.is_superuser:
            return error("No tenés acceso a esa escuela.", 403)
        request.session["active_school_id"] = school.id
        request.active_school = school
        request.school_membership = membership
        if connection.vendor == "postgresql":
            with connection.cursor() as cursor:
                cursor.execute("SELECT set_config('app.current_school_id', %s, true)", [str(school.pk)])
        if request.user.is_superuser:
            audit(request, "Cambio de escuela por plataforma", f"Acceso administrativo a {school.name}")
        return JsonResponse({"ok": True, "school_id": school.id})
    return error("Método no permitido.", 405)


@school_required
def api_dashboard(request):
    school = request.active_school
    role = active_role(request)
    allowed_audiences = notice_audiences(request)
    students = student_query(request)
    attendance = Attendance.objects.filter(school=school, date=timezone.localdate(), status=Attendance.Status.ABSENT)
    loans = Loan.objects.filter(school=school, returned_at__isnull=True)
    notices = Notice.objects.filter(school=school, audience__in=allowed_audiences)
    if role in {"alumno", "tutor"}:
        loans = loans.filter(student__in=students)
    return JsonResponse({"students": Student.objects.filter(school=school, status=Student.Status.ACTIVE).count(),
                         "absent": attendance.count(), "loans": loans.count(),
                         "notices": Notice.objects.filter(school=school).count(),
                         "recent": [serialize_notice(n) for n in notices.order_by("-id")[:5]]})


@school_required
def api_academic_followup(request):
    if request.method != "GET":
        return error("Método no permitido.", 405)
    role = active_role(request)
    if role not in ATTENDANCE_ROLES | GRADE_ROLES | {"alumno", "tutor"} and not request.user.is_superuser:
        return error("Tu perfil no puede consultar el seguimiento académico.", 403)
    school = request.active_school
    today = timezone.localdate()
    visible_students = student_query(request).filter(status=Student.Status.ACTIVE)
    student_ids = list(visible_students.values_list("pk", flat=True))
    absence_totals = dict(Attendance.objects.filter(school=school, student_id__in=student_ids,
        date__year=today.year, status=Attendance.Status.ABSENT).values("student_id")
        .annotate(days=Count("date", distinct=True)).values_list("student_id", "days"))
    students_by_id = {student.id: student for student in visible_students}
    absences = [{"student_id": student_id, "student": students_by_id[student_id].name, "days": days}
        for student_id, days in sorted(absence_totals.items(), key=lambda item: (-item[1], students_by_id[item[0]].name.casefold()))]

    active_period = AcademicPeriod.objects.filter(school=school, year=today.year,
        starts_on__lte=today, ends_on__gte=today).order_by("-starts_on", "order").first()
    pending = []
    if active_period and student_ids:
        enrollments = list(Enrollment.objects.filter(school=school, student_id__in=student_ids,
            academic_year=today.year, state=Enrollment.State.ACTIVE).select_related("student", "section"))
        section_ids = [enrollment.section_id for enrollment in enrollments]
        offerings = Offering.objects.filter(school=school, section_id__in=section_ids).select_related("subject")
        if role == "docente" and not request.user.is_superuser:
            offerings = offerings.filter(teacher_assignments__membership=request.school_membership)
        offerings_by_section = {}
        for offering in offerings.distinct():
            offerings_by_section.setdefault(offering.section_id, []).append(offering)
        existing = set(Grade.objects.filter(school=school, period=active_period,
            student_id__in=student_ids).values_list("student_id", "offering_id"))
        for enrollment in enrollments:
            for offering in offerings_by_section.get(enrollment.section_id, []):
                if (enrollment.student_id, offering.id) not in existing:
                    pending.append({"student_id": enrollment.student_id, "student": enrollment.student.name,
                        "subject": offering.subject.name, "period": active_period.name})
    return JsonResponse({"year": today.year, "period": active_period.name if active_period else None,
        "absence_count": len(absences), "absences": absences[:100],
        "pending_grade_count": len(pending), "pending_grades": pending[:100]})


ACADEMIC_ANALYTICS_ROLES = {"school_admin", "directivo", "secretaria", "preceptor"}


def _analytics_average(values):
    return round(float(sum(values) / len(values)), 2) if values else None


def _analytics_group(values, passing):
    approved = sum(value >= passing for value in values)
    return {"average": _analytics_average(values), "grade_count": len(values),
        "approved_count": approved,
        "approved_percentage": round(approved * 100 / len(values), 2) if values else None}


@school_required
def api_academic_analytics(request):
    if request.method != "GET":
        return error("Método no permitido.", 405)
    if active_role(request) not in ACADEMIC_ANALYTICS_ROLES:
        return error("Tu perfil no puede consultar las estadísticas académicas.", 403)
    school = request.active_school
    subscription = getattr(school, "subscription", None)
    if subscription:
        apply_due_subscription_plan(subscription)
    if (school.state != School.State.ACTIVE or not subscription
            or subscription.state != SchoolSubscription.State.ACTIVE
            or effective_subscription_plan(subscription) != SubscriptionPlan.PRO):
        return error("Las estadísticas académicas están disponibles en el plan Pro.", 403)

    try:
        year = int(request.GET.get("year", timezone.localdate().year))
        if not 1900 <= year <= 2100:
            raise ValueError
    except (TypeError, ValueError):
        return error("El ciclo lectivo no es válido.", 400)

    periods = AcademicPeriod.objects.filter(school=school, year=year).order_by("order", "name")
    sections = Section.objects.filter(school=school, academic_year=year).select_related("plan_year").order_by("plan_year__ordinal", "division", "shift")
    subjects = Subject.objects.filter(school=school).order_by("name")
    grades = Grade.objects.filter(school=school, student__school=school, period__school=school,
        period__year=year, offering__school=school, offering__section__school=school,
        offering__section__academic_year=year).select_related(
            "student", "period", "offering__section__plan_year", "offering__subject")
    selected = {"year": year, "period_id": None, "section_id": None, "subject_id": None}
    try:
        for key, model, queryset, field in (
            ("period_id", AcademicPeriod, periods, "period_id"),
            ("section_id", Section, sections, "offering__section_id"),
            ("subject_id", Subject, subjects, "offering__subject_id"),
        ):
            raw = request.GET.get(key)
            if raw:
                object_id = int(raw)
                if not queryset.filter(pk=object_id).exists():
                    return error(f"El filtro {key} no pertenece a este ciclo lectivo y escuela.", 400)
                selected[key] = object_id
                grades = grades.filter(**{field: object_id})
    except (TypeError, ValueError):
        return error("Uno de los filtros académicos no es válido.", 400)

    rows = list(grades.order_by("student__name", "period__order", "offering__subject__name"))
    scale = GradingScale.objects.filter(school=school).first() or GradingScale(school=school)
    passing = scale.passing
    values = [row.value for row in rows]
    student_groups, section_groups, subject_groups, period_groups = (defaultdict(list) for _ in range(4))
    student_sections = defaultdict(set)
    section_by_id = {}
    subject_by_id = {}
    period_by_id = {}
    distribution = defaultdict(int)
    for row in rows:
        section = row.offering.section
        subject = row.offering.subject
        period = row.period
        section_label = section.name or f"{section.plan_year.year_label} {section.division} · {section.get_shift_display()}"
        student_groups[row.student_id].append(row.value)
        section_groups[section.id].append(row.value)
        subject_groups[subject.id].append(row.value)
        period_groups[period.id].append(row.value)
        student_sections[row.student_id].add(section_label)
        section_by_id[section.id] = {"section_id": section.id, "section": section_label,
            "year_label": section.plan_year.year_label, "division": section.division,
            "shift": section.get_shift_display()}
        subject_by_id[subject.id] = {"subject_id": subject.id, "subject": subject.name,
            "kind": subject.kind}
        period_by_id[period.id] = {"period_id": period.id, "period": period.name,
            "order": period.order}
        distribution[row.value] += 1

    def grouped(groups, metadata, label_key, id_key):
        result = []
        for object_id, group_values in groups.items():
            row = dict(metadata[object_id])
            row.update(_analytics_group(group_values, passing))
            result.append(row)
        return sorted(result, key=lambda item: (str(item[label_key]).casefold(), item[id_key]))

    by_student = []
    students_by_id = {row.student_id: row.student for row in rows}
    for student_id, group_values in student_groups.items():
        item = {"student_id": student_id, "student": students_by_id[student_id].name,
            "sections": sorted(student_sections[student_id], key=str.casefold)}
        item.update(_analytics_group(group_values, passing))
        by_student.append(item)
    by_student.sort(key=lambda item: (item["student"].casefold(), item["student_id"]))

    years = list(AcademicPeriod.objects.filter(school=school).values_list("year", flat=True).distinct().order_by("-year"))
    if year not in years:
        years.insert(0, year)
    distribution_rows = [{"grade": float(value), "count": count}
        for value, count in sorted(distribution.items())]
    return JsonResponse({
        "year": year,
        "filters": selected,
        "scale": {"minimum": float(scale.minimum), "maximum": float(scale.maximum), "passing": float(passing)},
        "summary": _analytics_group(values, passing),
        "distribution": distribution_rows,
        "by_student": by_student,
        "by_section": grouped(section_groups, section_by_id, "section", "section_id"),
        "by_subject": grouped(subject_groups, subject_by_id, "subject", "subject_id"),
        "by_period": sorted((dict(period_by_id[period_id], **_analytics_group(group_values, passing))
            for period_id, group_values in period_groups.items()), key=lambda item: (item["order"], item["period"].casefold())),
        "options": {
            "years": years,
            "periods": [{"id": item.id, "name": item.name, "year": item.year, "order": item.order} for item in periods],
            "sections": [{"id": item.id, "label": item.name or f"{item.plan_year.year_label} {item.division} · {item.get_shift_display()}"} for item in sections],
            "subjects": [{"id": item.id, "name": item.name, "kind": item.kind} for item in subjects],
        },
    })


@school_required
def api_students(request):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        if role not in WRITE_ROLES | {"secretaria", "preceptor", "docente", "biblioteca", "alumno", "tutor"} and not request.user.is_superuser:
            return error("Tu perfil no puede consultar el padrón de alumnos.", 403)
        return JsonResponse([student_data(s, school) for s in student_query(request)], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, "school_admin", "directivo", "secretaria"):
        return error("Tu perfil no tiene permiso para crear alumnos.", 403)
    try:
        data = json_body(request)
        student, _ = create_student(school, data)
        audit(request, "Alumno", f"Alta de {student.name}")
        return JsonResponse(student_data(student, school), status=201)
    except (ValueError, KeyError, IntegrityError, Section.DoesNotExist) as exc:
        return error(exc or "No se pudo crear el alumno.")


def ensure_plan_year(school, year_label, school_type=None, orientation=""):
    school_type = school_type or school.school_type
    plan, _ = AcademicPlan.objects.get_or_create(school=school, name=f"Plan base {school.get_school_type_display()}",
        valid_from_year=timezone.localdate().year, defaults={"school_type": school_type, "jurisdiction": school.jurisdiction, "orientation": orientation})
    ordinal_match = re.search(r"\d+", year_label or "")
    ordinal = int(ordinal_match.group()) if ordinal_match else PlanYear.objects.filter(plan=plan).count() + 1
    py, _ = PlanYear.objects.get_or_create(school=school, plan=plan, ordinal=ordinal,
        defaults={"year_label": year_label or f"Año {ordinal}", "cycle": ""})
    return py


def ensure_section(school, year_label, division="A", shift="morning", orientation=""):
    py = ensure_plan_year(school, year_label, orientation=orientation)
    return Section.objects.get_or_create(school=school, plan_year=py, academic_year=timezone.localdate().year,
        division=division or "A", shift=shift or "morning", defaults={"name": f"{year_label} {division or 'A'}"})[0]


def create_student(school, data):
    name = str(data.get("name", "")).strip()
    if not name:
        raise ValueError("El nombre del alumno es obligatorio.")
    source_id = str(data.get("source_id", "")).strip()
    if source_id and Student.objects.filter(school=school, source_id=source_id).exists():
        raise ValueError(f"Ya existe un alumno con identificador {source_id}.")
    status = data.get("status", Student.Status.ACTIVE)
    if status not in Student.Status.values:
        status = Student.Status.ACTIVE
    email = str(data.get("email", "")).strip()
    if email:
        try:
            validate_email(email)
        except ValidationError:
            raise ValueError("El correo electrónico del alumno no es válido.")
    if data.get("section_id"):
        section = Section.objects.get(pk=data["section_id"], school=school)
    else:
        section = ensure_section(school, str(data.get("course", "")).strip(), str(data.get("division", "A")).strip(),
                                 str(data.get("shift", "morning")).strip(), str(data.get("orientation", "")).strip())
    student = Student.objects.create(school=school, source_id=source_id, name=name,
        email=email, status=status)
    Enrollment.objects.create(school=school, student=student, section=section, academic_year=section.academic_year)
    return student, section


@school_required
def api_attendance(request):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        if role not in ATTENDANCE_ROLES | {"alumno", "tutor", "docente"} and not request.user.is_superuser:
            return error("Tu perfil no tiene permiso para consultar asistencia.", 403)
        qs = Attendance.objects.filter(school=school).select_related("student")
        if role in {"alumno", "tutor"}:
            qs = qs.filter(student__in=student_query(request))
        elif role == "docente":
            qs = qs.filter(student__in=student_query(request)).filter(
                Q(offering__isnull=True) | Q(offering__teacher_assignments__membership=request.school_membership)).distinct()
        rows = [serialize_attendance(a) for a in qs.order_by("-date", "student__name")[:200]]
        return JsonResponse(rows, safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, *ATTENDANCE_ROLES, "docente"):
        return error("Tu perfil no tiene permiso para registrar asistencia.", 403)
    try:
        data = json_body(request)
        statuses = {"Presente": Attendance.Status.PRESENT, "Ausente": Attendance.Status.ABSENT,
                    "Tarde": Attendance.Status.LATE, "Justificado": Attendance.Status.EXCUSED,
                    **dict(Attendance.Status.choices)}
        records = data.get("records")
        if records is None:
            records = [data]
        if not isinstance(records, list) or not records or len(records) > 250:
            raise ValueError("La lista debe incluir entre 1 y 250 registros.")
        prepared = []
        seen = set()
        for row in records:
            if not isinstance(row, dict):
                raise ValueError("Cada registro de asistencia debe ser un objeto válido.")
            student = Student.objects.get(pk=int(row["student_id"]), school=school)
            status = statuses.get(row.get("status"))
            if not status:
                raise ValueError("Estado de asistencia inválido.")
            day = date.fromisoformat(row.get("date") or data.get("date"))
            offering_id = row.get("offering_id", data.get("offering_id"))
            offering = Offering.objects.get(pk=offering_id, school=school) if offering_id else None
            if role == "docente":
                if not offering:
                    raise ValueError("Elegí una materia o taller asignado para pasar asistencia.")
                if not TeacherAssignment.objects.filter(school=school, offering=offering, membership=request.school_membership).exists():
                    return error("No estás asignado a ese taller o materia.", 403)
            if offering and (enrollment_for(student, school) is None or enrollment_for(student, school).section_id != offering.section_id):
                raise ValueError("El alumno no pertenece a la comisión seleccionada.")
            key = (student.pk, day, offering.pk if offering else None)
            if key in seen:
                raise ValueError("La lista contiene más de un registro para el mismo alumno e instancia.")
            seen.add(key)
            prepared.append((student, day, status, offering, str(row.get("note", data.get("note", "")))[:500]))
        saved = []
        for student, day, status, offering, note in prepared:
            obj, _ = Attendance.objects.update_or_create(school=school, student=student, date=day, offering=offering,
                defaults={"status": status, "note": note, "recorded_by": request.user})
            saved.append(obj)
        audit(request, "Asistencia", f"{len(saved)} registros · {saved[0].date.isoformat()}")
        if data.get("records") is not None:
            return JsonResponse({"ok": True, "records": [serialize_attendance(obj) for obj in saved]}, status=201)
        return JsonResponse(serialize_attendance(saved[0]), status=201)
    except (ValueError, TypeError, KeyError, Student.DoesNotExist, Offering.DoesNotExist, IntegrityError) as exc:
        return error(exc or "No se pudo registrar la asistencia.")


@school_required
def api_grades(request):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        if role not in GRADE_ROLES | {"alumno", "tutor"} and not request.user.is_superuser:
            return error("Tu perfil no puede consultar calificaciones.", 403)
        qs = Grade.objects.filter(school=school).select_related("student", "offering__subject", "period")
        if role in {"alumno", "tutor", "docente"}:
            qs = qs.filter(student__in=student_query(request))
        if role == "docente" and not request.user.is_superuser:
            qs = qs.filter(offering__teacher_assignments__membership=request.school_membership)
        return JsonResponse([serialize_grade(g) for g in qs.order_by("period__year", "period__order", "student__name", "offering__subject__name")], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, *GRADE_ROLES):
        return error("Tu perfil no tiene permiso para registrar calificaciones.", 403)
    try:
        data = json_body(request)
        student = Student.objects.get(pk=int(data["student_id"]), school=school)
        if role == "docente" and not can_see_student(request, student):
            return error("No estás asignado al curso de ese alumno.", 403)
        enrollment = enrollment_for(student, school)
        if not enrollment:
            raise ValueError("El alumno no tiene una inscripción activa este ciclo lectivo.")
        subject = Subject.objects.get(school=school, name=data["subject"])
        if data.get("period_id"):
            period = AcademicPeriod.objects.get(school=school, pk=data["period_id"])
        else:
            period = AcademicPeriod.objects.get(school=school, name=data["period"], year=int(data.get("year", timezone.localdate().year)))
        scale = GradingScale.objects.filter(school=school).first()
        value = float(data["grade"])
        if scale and not float(scale.minimum) <= value <= float(scale.maximum):
            raise ValueError(f"La calificación debe estar entre {scale.minimum:g} y {scale.maximum:g}.")
        offering, _ = Offering.objects.get_or_create(school=school, section=enrollment.section, subject=subject)
        if role == "docente" and not TeacherAssignment.objects.filter(school=school, offering=offering, membership=request.school_membership).exists():
            return error("No estás asignado a esa materia.", 403)
        grade, _ = Grade.objects.update_or_create(school=school, student=student, offering=offering, period=period,
            defaults={"value": value, "note": str(data.get("note", ""))[:500], "recorded_by": request.user})
        audit(request, "Calificación", f"{subject.name} · {student.name} · {period.name}")
        return JsonResponse(serialize_grade(grade), status=201)
    except (ValueError, KeyError, Student.DoesNotExist, Subject.DoesNotExist, AcademicPeriod.DoesNotExist, IntegrityError) as exc:
        return error(exc or "No se pudo registrar la calificación.")


REPORT_CARD_SLOTS = [
    (AcademicPeriod.ReportSlot.FIRST, "1° informe"),
    (AcademicPeriod.ReportSlot.SECOND, "2° informe"),
    (AcademicPeriod.ReportSlot.THIRD, "3° informe"),
    (AcademicPeriod.ReportSlot.EXTENDED, "Período extendido"),
    (AcademicPeriod.ReportSlot.FINAL, "Informe final"),
]


def report_slot_for(period):
    if period.report_slot != AcademicPeriod.ReportSlot.AUTO:
        return period.report_slot
    name = unicodedata.normalize("NFKD", period.name).encode("ascii", "ignore").decode().lower()
    if any(token in name for token in ("extend", "recuper", "intens", "previo")):
        return AcademicPeriod.ReportSlot.EXTENDED
    if any(token in name for token in ("final", "anual", "cierre")):
        return AcademicPeriod.ReportSlot.FINAL
    return {1: AcademicPeriod.ReportSlot.FIRST, 2: AcademicPeriod.ReportSlot.SECOND,
            3: AcademicPeriod.ReportSlot.THIRD}.get(period.order, AcademicPeriod.ReportSlot.FINAL)


def report_card_access(request):
    return active_role(request) in GRADE_ROLES | {"alumno", "tutor"} or request.user.is_superuser


@school_required
def api_report_card_options(request):
    if request.method != "GET":
        return error("Método no permitido.", 405)
    if not report_card_access(request):
        return error("Tu perfil no puede consultar boletines.", 403)
    school = request.active_school
    visible_students = student_query(request).values_list("pk", flat=True)
    enrollments = Enrollment.objects.filter(school=school, student_id__in=visible_students).select_related(
        "student", "section__plan_year"
    ).order_by("student__name", "-academic_year")
    students = {}
    for enrollment in enrollments:
        row = students.setdefault(enrollment.student_id, {"id": enrollment.student_id, "name": enrollment.student.name, "enrollments": []})
        row["enrollments"].append({
            "year": enrollment.academic_year,
            "course": enrollment.section.plan_year.year_label,
            "division": enrollment.section.division,
            "section": enrollment.section.name or f"{enrollment.section.plan_year.year_label} {enrollment.section.division}",
            "shift": enrollment.section.get_shift_display(),
        })
    return JsonResponse({"students": list(students.values())})


def build_report_card(request, student, enrollment):
    school = request.active_school
    role = active_role(request)
    periods = list(AcademicPeriod.objects.filter(school=school, year=enrollment.academic_year).order_by("order", "name"))
    period_by_slot = {}
    for period in periods:
        period_by_slot[report_slot_for(period)] = period

    offerings = list(Offering.objects.filter(school=school, section=enrollment.section).select_related("subject").order_by("subject__name"))
    if role == "docente" and not request.user.is_superuser:
        offerings = [offering for offering in offerings if offering.teacher_assignments.filter(membership=request.school_membership).exists()]
    subjects = {entry.subject_id: entry.subject.name for entry in enrollment.section.plan_year.subject_entries.select_related("subject").order_by("subject__name")}
    for offering in offerings:
        subjects.setdefault(offering.subject_id, offering.subject.name)

    grades = Grade.objects.filter(
        school=school, student=student, offering__section=enrollment.section, period__year=enrollment.academic_year,
    ).select_related("offering__subject", "period").order_by("updated_at", "pk")
    if role == "docente" and not request.user.is_superuser:
        grades = grades.filter(offering__teacher_assignments__membership=request.school_membership)
    grade_values = {}
    for grade in grades:
        slot = report_slot_for(grade.period)
        selected_period = period_by_slot.get(slot)
        if selected_period and selected_period.pk == grade.period_id:
            grade_values.setdefault(grade.offering.subject_id, {})[slot] = format(grade.value.normalize(), "f")

    slots = []
    for code, label in REPORT_CARD_SLOTS:
        period = period_by_slot.get(code)
        attendance = {"school_days": None, "absences": None, "configured": bool(period and period.starts_on and period.ends_on)}
        if attendance["configured"]:
            enrolled_ids = Enrollment.objects.filter(
                school=school, section=enrollment.section, academic_year=enrollment.academic_year
            ).values_list("student_id", flat=True)
            daily_records = Attendance.objects.filter(
                school=school, student_id__in=enrolled_ids, offering__isnull=True,
                date__gte=period.starts_on, date__lte=period.ends_on,
            )
            daily_dates = set(daily_records.values_list("date", flat=True).distinct())
            class_records = Attendance.objects.filter(
                school=school, student_id__in=enrolled_ids, offering__section=enrollment.section,
                date__gte=period.starts_on, date__lte=period.ends_on,
            )
            class_dates = set(class_records.values_list("date", flat=True).distinct())
            attendance["school_days"] = len(daily_dates | class_dates)
            student_daily = dict(daily_records.filter(student=student).values_list("date", "status"))
            daily_absences = {day for day, status in student_daily.items() if status == Attendance.Status.ABSENT}
            class_absences = set(class_records.filter(student=student, status=Attendance.Status.ABSENT).values_list("date", flat=True).distinct())
            attendance["absences"] = len(daily_absences | (class_absences - set(student_daily)))
        slots.append({
            "code": code, "label": label, "period": period.name if period else "",
            "starts_on": period.starts_on.isoformat() if period and period.starts_on else None,
            "ends_on": period.ends_on.isoformat() if period and period.ends_on else None,
            "attendance": attendance,
        })

    return {
        "school": school.name,
        "student": student.name,
        "student_id": student.pk,
        "source_id": student.source_id,
        "academic_year": enrollment.academic_year,
        "course": enrollment.section.plan_year.year_label,
        "division": enrollment.section.division,
        "section": enrollment.section.name or f"{enrollment.section.plan_year.year_label} {enrollment.section.division}",
        "shift": enrollment.section.get_shift_display(),
        "jurisdiction": school.jurisdiction,
        "slots": slots,
        "subjects": [
            {"name": name, "grades": [grade_values.get(subject_id, {}).get(code, "") for code, _ in REPORT_CARD_SLOTS]}
            for subject_id, name in sorted(subjects.items(), key=lambda item: item[1].casefold())
        ],
        "generated_on": timezone.localdate().isoformat(),
    }


def render_report_card_pdf(data):
    import io
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Boletín - {data['student']}", author=data["school"])
    base = getSampleStyleSheet()
    school_style = ParagraphStyle("ReportSchool", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=colors.HexColor("#263942"))
    title_style = ParagraphStyle("ReportTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=15, leading=18, alignment=1, textColor=colors.HexColor("#263942"), spaceAfter=4)
    label_style = ParagraphStyle("ReportLabel", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7, leading=8, textColor=colors.HexColor("#39464b"))
    cell_style = ParagraphStyle("ReportCell", parent=base["Normal"], fontName="Helvetica", fontSize=7.1, leading=8.4, textColor=colors.HexColor("#263238"))
    center_style = ParagraphStyle("ReportCenter", parent=cell_style, alignment=1)
    header_style = ParagraphStyle("ReportHeader", parent=label_style, fontSize=6.6, leading=7.4, textColor=colors.white, alignment=1)
    meta_style = ParagraphStyle("ReportMeta", parent=cell_style, fontSize=8, leading=10)
    story = [Paragraph(escape(data["school"]), school_style), Spacer(1, 2 * mm), Paragraph("BOLETÍN DE CALIFICACIONES", title_style)]
    width = A4[0] - 28 * mm
    thirds = width / 3
    metadata = Table([[
        Paragraph(f"CICLO LECTIVO <b>{data['academic_year']}</b>", meta_style),
        Paragraph(f"AÑO <b>{escape(data['course'])}</b>", meta_style),
        Paragraph(f"SECCIÓN <b>{escape(data['division'])}</b>", meta_style),
    ]], colWidths=[thirds] * 3, rowHeights=[9 * mm])
    metadata.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eef1f1")), ("BOX", (0, 0), (-1, -1), .6, colors.HexColor("#9da8aa")), ("INNERGRID", (0, 0), (-1, -1), .4, colors.HexColor("#bac2c3")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    story.extend([metadata, Spacer(1, 2 * mm)])
    student_line = Table([[
        Paragraph(f"ALUMNO/A: <b>{escape(data['student'])}</b>", meta_style),
        Paragraph(f"COMISIÓN: <b>{escape(data['section'])}</b>", meta_style),
    ]], colWidths=[width * .62, width * .38], rowHeights=[8 * mm])
    student_line.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .6, colors.HexColor("#9da8aa")), ("INNERGRID", (0, 0), (-1, -1), .4, colors.HexColor("#bac2c3")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    story.extend([student_line, Spacer(1, 3 * mm)])

    column_widths = [width * .28] + [width * .144] * 5
    grid = [[Paragraph("ÁREA CURRICULAR", header_style)] + [Paragraph(label.upper(), header_style) for _, label in REPORT_CARD_SLOTS]]
    for subject in data["subjects"]:
        grid.append([Paragraph(escape(subject["name"]), cell_style)] + [Paragraph(escape(value), center_style) if value else "" for value in subject["grades"]])
    for label, key in (("DÍAS HÁBILES", "school_days"), ("INASISTENCIAS", "absences")):
        grid.append([Paragraph(label, label_style)] + [str(slot["attendance"][key]) if slot["attendance"][key] is not None else "—" for slot in data["slots"]])
    marks = Table(grid, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
    marks.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#36454b")), ("BOX", (0, 0), (-1, -1), .7, colors.HexColor("#657276")), ("INNERGRID", (0, 0), (-1, -1), .45, colors.HexColor("#8d989a")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 1), (-1, -1), "CENTER"), ("LEFTPADDING", (0, 0), (0, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("BACKGROUND", (0, -2), (-1, -1), colors.HexColor("#f0f2f2"))]))
    story.extend([marks, Spacer(1, 2 * mm)])
    if any(not slot["attendance"]["configured"] for slot in data["slots"]):
        note_style = ParagraphStyle("ReportNote", parent=cell_style, fontSize=6.5, textColor=colors.HexColor("#697579"))
        story.extend([Paragraph("Las casillas de asistencia con guion requieren configurar las fechas de inicio y cierre del período.", note_style), Spacer(1, 2 * mm)])

    signatures = [
        [Paragraph("FIRMA DEL/DE LA DOCENTE<br/><br/><br/>Aclaración: __________________________", cell_style), Paragraph("FIRMA DEL/DE LA DIRECTOR/A<br/><br/><br/>Aclaración: __________________________", cell_style)],
        [Paragraph("FIRMA DEL/DE LA ESTUDIANTE<br/><br/><br/>Aclaración: __________________________", cell_style), Paragraph("FIRMA DEL ADULTO RESPONSABLE<br/><br/><br/>Aclaración: __________________________", cell_style)],
    ]
    signatures_table = Table(signatures, colWidths=[width / 2] * 2, rowHeights=[21 * mm, 21 * mm])
    signatures_table.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .7, colors.HexColor("#657276")), ("INNERGRID", (0, 0), (-1, -1), .45, colors.HexColor("#8d989a")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
    story.extend([signatures_table, Spacer(1, 3 * mm)])
    footer = Table([[
        Paragraph("FIRMA DE LA DIRECCIÓN<br/><br/>__________________________________", cell_style),
        Paragraph("PROMUEVE A: ____________________<br/><br/>CONTINÚA EN: ____________________", cell_style),
    ]], colWidths=[width / 2] * 2, rowHeights=[17 * mm])
    footer.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3)]))
    story.append(footer)
    doc.build(story)
    return buffer.getvalue()


@school_required
def api_report_card(request, student_id):
    if request.method != "GET":
        return error("Método no permitido.", 405)
    if not report_card_access(request):
        return error("Tu perfil no puede consultar boletines.", 403)
    student = student_query(request).filter(pk=student_id).first()
    if not student:
        return error("No se encontró el alumno o no tenés acceso a su boletín.", 404)
    try:
        year = int(request.GET.get("year", timezone.localdate().year))
    except (TypeError, ValueError):
        return error("El ciclo lectivo no es válido.")
    enrollment = Enrollment.objects.select_related("section__plan_year").filter(
        school=request.active_school, student=student, academic_year=year
    ).first()
    if not enrollment:
        return error("El alumno no tiene una inscripción en ese ciclo lectivo.", 404)
    data = build_report_card(request, student, enrollment)
    if request.GET.get("format") == "pdf":
        response = HttpResponse(render_report_card_pdf(data), content_type="application/pdf")
        ascii_name = unicodedata.normalize("NFKD", student.name).encode("ascii", "ignore").decode()
        safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "-", ascii_name).strip("-").lower() or "alumno"
        response["Content-Disposition"] = f'attachment; filename="boletin-{safe_name}-{year}.pdf"'
        response["Cache-Control"] = "no-store"
        return response
    return JsonResponse(data)


@school_required
def api_books(request):
    school = request.active_school
    if request.method == "GET":
        if not allows(request, *LIBRARY_ROLES, "alumno", "tutor"):
            return error("Tu perfil no puede consultar el catálogo.", 403)
        return JsonResponse([serialize_book(b) for b in Book.objects.filter(school=school)], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, *LIBRARY_ROLES):
        return error("Tu perfil no puede administrar el catálogo.", 403)
    try:
        data = json_body(request)
        obj = Book.objects.create(school=school, title=str(data["title"]).strip(), author=str(data["author"]).strip(),
                                  code=str(data["code"]).strip(), copies=max(1, int(data.get("copies", 1))))
        audit(request, "Catálogo", f"Alta de libro: {obj.title}")
        return JsonResponse(serialize_book(obj), status=201)
    except (KeyError, ValueError, IntegrityError) as exc:
        return error(exc or "Revisá los datos del libro; el código debe ser único en la escuela.")


@school_required
def api_loans(request, object_id=None):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        if role not in LIBRARY_ROLES | {"alumno", "tutor"} and not request.user.is_superuser:
            return error("Tu perfil no puede consultar préstamos.", 403)
        qs = Loan.objects.filter(school=school).select_related("book", "student")
        if role in {"alumno", "tutor"}:
            qs = qs.filter(student__in=student_query(request))
        return JsonResponse([serialize_loan(l) for l in qs], safe=False)
    if request.method == "POST":
        if not allows(request, *LIBRARY_ROLES):
            return error("Tu perfil no puede registrar préstamos.", 403)
        try:
            data = json_body(request)
            book = Book.objects.select_for_update().get(pk=int(data["book_id"]), school=school)
            student = Student.objects.get(pk=int(data["student_id"]), school=school)
            available = book.copies - book.loans.filter(returned_at__isnull=True).count()
            if available < 1:
                raise ValueError("No hay ejemplares disponibles para prestar.")
            due = date.fromisoformat(data["due_at"])
            if due < timezone.localdate():
                raise ValueError("La fecha de devolución no puede ser anterior a hoy.")
            borrowed = timezone.now()
            if data.get("borrowed_at"):
                borrowed = timezone.make_aware(datetime.fromisoformat(data["borrowed_at"])) if timezone.is_naive(datetime.fromisoformat(data["borrowed_at"])) else datetime.fromisoformat(data["borrowed_at"])
            loan = Loan.objects.create(school=school, book=book, student=student, due_at=due, borrowed_at=borrowed, recorded_by=request.user)
            audit(request, "Préstamo", f"{book.title} → {student.name}")
            return JsonResponse(serialize_loan(loan), status=201)
        except (KeyError, ValueError, Book.DoesNotExist, Student.DoesNotExist, IntegrityError) as exc:
            return error(exc or "No se pudo registrar el préstamo.")
    if request.method == "PUT" and object_id:
        if not allows(request, *LIBRARY_ROLES):
            return error("Tu perfil no puede registrar devoluciones.", 403)
        loan = Loan.objects.filter(pk=object_id, school=school, returned_at__isnull=True).select_related("book", "student").first()
        if not loan:
            return error("No existe un préstamo activo con ese identificador.", 404)
        loan.returned_at = timezone.now()
        loan.save(update_fields=("returned_at",))
        audit(request, "Devolución", f"{loan.book.title} · {loan.student.name}")
        return JsonResponse(serialize_loan(loan))
    return error("Método no permitido.", 405)


@school_required
def api_lost(request):
    school = request.active_school
    if request.method == "GET":
        return JsonResponse([serialize_lost(x) for x in LostItem.objects.filter(school=school).order_by("-id")], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, *NOTICE_ROLES):
        return error("Solo el personal autorizado puede publicar objetos encontrados.", 403)
    try:
        data = json_body(request)
        item = LostItem.objects.create(school=school, title=str(data["title"]).strip(), description=str(data.get("description", "")),
            place=str(data.get("place", "")), created_by=request.user)
        audit(request, "Objeto perdido", f"Publicado: {item.title}")
        return JsonResponse(serialize_lost(item), status=201)
    except (KeyError, ValueError) as exc:
        return error(exc)


@school_required
def api_claims(request, object_id=None):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        qs = Claim.objects.filter(school=school).select_related("item", "student")
        if role == "alumno": qs = qs.filter(student__account=request.user)
        elif role == "tutor": qs = qs.filter(student__guardian_links__guardian=request.user)
        elif not allows(request, *NOTICE_ROLES): return error("Tu perfil no puede consultar reclamos.", 403)
        return JsonResponse([serialize_claim(c) for c in qs.order_by("-id")], safe=False)
    if request.method == "POST":
        if role != "alumno":
            return error("Solo un alumno puede enviar un reclamo.", 403)
        student = Student.objects.filter(school=school, account=request.user).first()
        if not student:
            return error("La cuenta no está vinculada a un alumno.", 400)
        try:
            data = json_body(request)
            item = LostItem.objects.get(pk=int(data["item_id"]), school=school, status=LostItem.State.PUBLISHED)
            claim = Claim.objects.create(school=school, item=item, student=student, message=str(data.get("message", ""))[:1000])
            item.status = LostItem.State.CLAIMED
            item.save(update_fields=("status",))
            audit(request, "Reclamo", f"{student.name} reclamó {item.title}")
            return JsonResponse(serialize_claim(claim), status=201)
        except (KeyError, ValueError, LostItem.DoesNotExist) as exc:
            return error(exc or "El objeto no está disponible para reclamo.")
    if request.method == "PUT" and object_id:
        if not allows(request, *NOTICE_ROLES): return error("Tu perfil no puede resolver reclamos.", 403)
        try:
            data = json_body(request)
            status = {"Resuelto": Claim.State.RESOLVED, "Rechazado": Claim.State.REJECTED,
                      **dict(Claim.State.choices)}.get(data.get("status", "Resuelto"))
            if not status: raise ValueError("Estado de reclamo inválido.")
            claim = Claim.objects.get(pk=object_id, school=school)
            claim.status = status
            claim.save(update_fields=("status",))
            if status == Claim.State.RESOLVED:
                claim.item.status = LostItem.State.RETURNED
                claim.item.save(update_fields=("status",))
            audit(request, "Reclamo", f"{claim.item.title}: {claim.get_status_display()}")
            return JsonResponse(serialize_claim(claim))
        except (ValueError, Claim.DoesNotExist) as exc:
            return error(exc)
    return error("Método no permitido.", 405)


@school_required
def api_notices(request):
    school = request.active_school
    if request.method == "GET":
        qs = Notice.objects.filter(school=school, audience__in=notice_audiences(request)).prefetch_related("reads")
        return JsonResponse([serialize_notice(n, request.user) for n in qs], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if not allows(request, *NOTICE_ROLES):
        return error("Tu perfil no puede publicar avisos.", 403)
    try:
        data = json_body(request)
        audience = audience_code(data.get("audience", "Todos"))
        notice = Notice.objects.create(school=school, title=str(data["title"]).strip(), body=str(data["body"]).strip(),
                                       audience=audience, created_by=request.user)
        audit(request, "Aviso", notice.title)
        return JsonResponse(serialize_notice(notice), status=201)
    except (KeyError, ValueError) as exc:
        return error(exc)


@school_required
def api_notice_read(request, notice_id):
    if request.method != "POST":
        return error("Método no permitido.", 405)
    if active_role(request) not in {"alumno", "tutor"}:
        return error("La confirmación de lectura está disponible para alumnos y familias.", 403)
    notice = Notice.objects.filter(school=request.active_school,
        audience__in=notice_audiences(request), pk=notice_id).first()
    if not notice:
        return error("No se encontró el aviso.", 404)
    receipt, created = NoticeRead.objects.get_or_create(school=request.active_school,
        notice=notice, user=request.user)
    return JsonResponse({"notice_id": notice.id, "is_read": True,
        "read_at": receipt.read_at.isoformat()}, status=201 if created else 200)


def parse_event_datetime(value, label):
    try:
        value = datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        raise ValueError(f"{label}: ingresá una fecha y hora válidas.")
    if timezone.is_naive(value):
        value = timezone.make_aware(value, timezone.get_current_timezone())
    return value


@school_required
def api_events(request, event_id=None):
    school = request.active_school
    role = active_role(request)
    if request.method == "GET":
        audience = set(Notice.Audience.values) if role in NOTICE_ROLES or request.user.is_superuser else notice_audiences(request)
        events = SchoolEvent.objects.filter(school=school, audience__in=audience,
            starts_at__gte=timezone.now() - timedelta(days=30)).order_by("starts_at")[:300]
        return JsonResponse([serialize_event(event) for event in events], safe=False)
    if not allows(request, *NOTICE_ROLES):
        return error("Tu perfil no puede administrar el calendario.", 403)
    if request.method == "POST" and event_id is None:
        try:
            data = json_body(request)
            title = str(data.get("title", "")).strip()
            if not title:
                raise ValueError("Ingresá el nombre del evento.")
            starts_at = parse_event_datetime(data.get("starts_at"), "Inicio")
            ends_at = parse_event_datetime(data.get("ends_at"), "Fin") if data.get("ends_at") else None
            if ends_at and ends_at < starts_at:
                raise ValueError("El cierre no puede ser anterior al inicio.")
            audience = audience_code(data.get("audience", "Todos"))
            event = SchoolEvent.objects.create(school=school, title=title[:180],
                description=str(data.get("description", "")).strip(), starts_at=starts_at,
                ends_at=ends_at, audience=audience, created_by=request.user)
            audit(request, "Calendario", f"Evento: {event.title}")
            return JsonResponse(serialize_event(event), status=201)
        except (ValueError, KeyError) as exc:
            return error(exc)
    if request.method == "DELETE" and event_id is not None:
        event = SchoolEvent.objects.filter(school=school, pk=event_id).first()
        if not event:
            return error("No se encontró el evento.", 404)
        audit(request, "Calendario", f"Eliminó evento: {event.title}")
        event.delete()
        return JsonResponse({"ok": True})
    return error("Método no permitido.", 405)


@school_required
def api_academic_options(request):
    school = request.active_school
    role = active_role(request)
    if role not in ACADEMIC_ROLES | {"alumno", "tutor"} and not request.user.is_superuser:
        return error("Tu perfil no puede consultar opciones académicas.", 403)
    scale = GradingScale.objects.filter(school=school).first()
    subjects = Subject.objects.filter(school=school, is_active=True)
    if role == "docente" and not request.user.is_superuser:
        subjects = subjects.filter(offerings__teacher_assignments__membership=request.school_membership).distinct()
    return JsonResponse({"courses": [{"id": c.id, "name": c.name} for c in Course.objects.filter(school=school, is_active=True)],
        "subjects": [{"id": s.id, "name": s.name, "kind": s.kind} for s in subjects],
        "periods": [{"id": p.id, "name": p.name, "year": p.year, "order": p.order,
                      "report_slot": p.report_slot, "starts_on": p.starts_on.isoformat() if p.starts_on else "",
                      "ends_on": p.ends_on.isoformat() if p.ends_on else ""}
                     for p in AcademicPeriod.objects.filter(school=school)],
        "scale": {"min": float(scale.minimum), "max": float(scale.maximum), "passing": float(scale.passing)} if scale else {"min": 1, "max": 10, "passing": 6}})


def send_invite(request, user):
    uid = urlsafe_base64_encode(force_bytes(user.pk))
    token = default_token_generator.make_token(user)
    link = request.build_absolute_uri(reverse("password_reset_confirm", kwargs={"uidb64": uid, "token": token}))
    send_mail("Activá tu cuenta de Nexo Escolar", f"Hola {user.name},\n\nPara definir tu contraseña y activar tu cuenta, abrí este enlace (válido por 24 horas):\n{link}\n", settings.DEFAULT_FROM_EMAIL, [user.email], fail_silently=False)


def send_subscription_status_email(subscription, status, charge=None):
    school_name = subscription.school.name
    contact_name = subscription.contact_name
    if status == SchoolSubscription.State.CANCELED:
        subject = "La suscripción de tu escuela fue cancelada"
        body = (
            f"Hola {contact_name},\n\n"
            f"Te informamos que la suscripción de {school_name} fue cancelada el "
            f"{timezone.localdate():%d/%m/%Y}. El acceso a Nexo Escolar quedó suspendido.\n\n"
            "Los datos de la escuela permanecen guardados. Para solicitar la reactivación, "
            "contactá a la administración de Nexo Escolar.\n"
        )
    elif status == SchoolSubscription.State.ACTIVE:
        if not charge:
            raise ValueError("Falta el cargo mensual para notificar la reactivación.")
        payment_status = "pagado" if charge.state == SubscriptionCharge.State.PAID else "pendiente"
        period = charge.period_start.strftime("%m/%Y") if charge.period_start else timezone.localdate().strftime("%m/%Y")
        subject = "La suscripción de tu escuela fue reactivada"
        body = (
            f"Hola {contact_name},\n\n"
            f"La suscripción de {school_name} al plan {subscription.get_plan_display()} fue reactivada. "
            "El acceso a Nexo Escolar ya está habilitado.\n\n"
            f"El abono mensual de {period}, por ${charge.amount_ars} ARS, figura como {payment_status}.\n"
        )
    elif status == SchoolSubscription.State.PENDING:
        subject = "El alta de tu escuela fue reabierta"
        body = (
            f"Hola {contact_name},\n\n"
            f"El proceso de alta de {school_name} fue reabierto. La escuela sigue pendiente de pago; "
            "el acceso se habilitará cuando se confirmen el alta y el primer abono.\n"
        )
    else:
        raise ValueError("Estado de suscripción no válido para enviar un aviso.")
    sent = send_mail(subject, body, settings.DEFAULT_FROM_EMAIL, [subscription.contact_email], fail_silently=False)
    if sent != 1:
        raise RuntimeError("El servidor de correo no aceptó el mensaje.")
    return sent


@transaction.atomic
def create_school_member(request, school, data):
    email = str(data.get("email", "")).strip().lower()
    name = str(data.get("name", "")).strip()
    role = ROLE_ALIASES.get(data.get("role"))
    if not email or "@" not in email or not name:
        raise ValueError("Ingresá nombre y correo electrónico válidos.")
    try:
        validate_email(email)
    except ValidationError:
        raise ValueError("El correo electrónico ingresado no es válido.")
    if role not in Membership.Role.values:
        raise ValueError("Rol inválido.")
    linked_ids = data.get("linked_student_ids") or ([data.get("linked_student_id")] if data.get("linked_student_id") else [])
    linked_students = []
    if role == Membership.Role.STUDENT:
        if len(linked_ids) != 1:
            raise ValueError("Vinculá la cuenta del alumno con exactamente un perfil.")
    elif role == Membership.Role.GUARDIAN and not linked_ids:
        raise ValueError("Vinculá el tutor con al menos un alumno.")
    if linked_ids:
        try:
            student_ids = [int(value) for value in linked_ids]
        except (TypeError, ValueError):
            raise ValueError("La selección de alumnos no es válida.")
        linked_students = list(Student.objects.filter(school=school, pk__in=student_ids))
        if len(linked_students) != len(set(student_ids)):
            raise ValueError("Uno o más alumnos no pertenecen a esta escuela.")
    user, created = User.objects.get_or_create(email=email, defaults={"name": name, "is_active": True})
    if not user.is_active:
        raise ValueError("La cuenta global está desactivada. Un administrador debe reactivarla antes de invitarla.")
    if Membership.objects.filter(school=school, user=user).exists():
        raise ValueError("Ya existe una cuenta de esa persona en esta escuela.")
    if role == Membership.Role.STUDENT and linked_students[0].account_id and linked_students[0].account_id != user.id:
        raise ValueError("Ese alumno ya tiene otra cuenta vinculada.")
    if created:
        user.set_unusable_password()
        user.save(update_fields=("password",))
    Membership.objects.create(school=school, user=user, role=role)
    if role == Membership.Role.STUDENT:
        linked_students[0].account = user
        linked_students[0].save(update_fields=("account",))
    elif role == Membership.Role.GUARDIAN:
        for student in linked_students:
            StudentGuardian.objects.get_or_create(school=school, student=student, guardian=user)
    send_invite(request, user)
    audit(request, "Cuenta escolar", f"Invitación para {email} · {role}")
    return user


@school_required
def api_settings(request):
    school = request.active_school
    if not allows(request, "school_admin", "directivo"):
        return error("Solo dirección puede administrar la configuración escolar.", 403)
    if request.method == "GET":
        scale = GradingScale.objects.filter(school=school).first()
        users = Membership.objects.filter(school=school).select_related("user").order_by("user__name")
        return JsonResponse({"courses": [{"id": x.id, "name": x.name} for x in Course.objects.filter(school=school)],
            "subjects": [{"id": x.id, "name": x.name, "kind": x.kind} for x in Subject.objects.filter(school=school)],
            "periods": [{"id": x.id, "name": x.name, "year": x.year, "order": x.order,
                         "report_slot": x.report_slot, "starts_on": x.starts_on.isoformat() if x.starts_on else "",
                         "ends_on": x.ends_on.isoformat() if x.ends_on else ""}
                        for x in AcademicPeriod.objects.filter(school=school)],
            "scale": {"min": float(scale.minimum), "max": float(scale.maximum), "passing": float(scale.passing)} if scale else {"min": 1, "max": 10, "passing": 6},
        "users": [{"id": m.user_id, "membership_id": m.id, "name": m.user.name, "email": m.user.email, "role": m.role} for m in users],
            "plans": [serialize_plan(p) for p in AcademicPlan.objects.filter(school=school)],
            "theme": school.theme})
    if request.method != "POST": return error("Método no permitido.", 405)
    try:
        data = json_body(request)
        kind = data.get("kind")
        if kind == "course":
            obj = Course.objects.create(school=school, name=str(data["name"]).strip())
            name = obj.name
        elif kind == "subject":
            obj = Subject.objects.create(school=school, name=str(data["name"]).strip(), kind=data.get("subject_kind", Subject.Kind.SUBJECT))
            name = obj.name
        elif kind == "period":
            report_slot = data.get("report_slot", AcademicPeriod.ReportSlot.AUTO)
            if report_slot not in AcademicPeriod.ReportSlot.values:
                raise ValueError("Elegí una columna de boletín válida.")
            starts_on = date.fromisoformat(data["starts_on"]) if data.get("starts_on") else None
            ends_on = date.fromisoformat(data["ends_on"]) if data.get("ends_on") else None
            if bool(starts_on) != bool(ends_on):
                raise ValueError("Cargá las dos fechas del período o dejalas vacías.")
            if starts_on and ends_on and starts_on > ends_on:
                raise ValueError("La fecha de inicio no puede ser posterior a la fecha de cierre.")
            defaults = {
                "name": str(data["name"]).strip(), "year": int(data["year"]),
                "order": int(data.get("order", 1)), "report_slot": report_slot,
                "starts_on": starts_on, "ends_on": ends_on,
            }
            if data.get("period_id"):
                obj = AcademicPeriod.objects.get(school=school, pk=int(data["period_id"]))
                for key, value in defaults.items():
                    setattr(obj, key, value)
                obj.save()
            else:
                obj = AcademicPeriod.objects.create(school=school, **defaults)
            name = obj.name
        elif kind == "scale":
            low, high = float(data["min"]), float(data["max"])
            if low >= high: raise ValueError("La nota máxima debe ser mayor que la mínima.")
            obj, _ = GradingScale.objects.update_or_create(school=school, name="Escala general", defaults={"minimum": low, "maximum": high, "passing": float(data.get("passing", low))})
            name = obj.name
        elif kind == "school_theme":
            themes = dict(School.Theme.choices)
            theme = data.get("theme")
            if theme not in themes: raise ValueError("Elegí un tema institucional disponible.")
            school.theme = theme
            school.save(update_fields=("theme",))
            name = themes[theme]
        elif kind == "user":
            user = create_school_member(request, school, data)
            name = user.name
        else:
            raise ValueError("Tipo de configuración no válido.")
        if kind != "user": audit(request, "Configuración", f"Alta o cambio de {kind}: {name}")
        return JsonResponse({"ok": True, "name": name}, status=201)
    except (KeyError, ValueError, TypeError, IntegrityError, Student.DoesNotExist, School.DoesNotExist, AcademicPeriod.DoesNotExist) as exc:
        return error(exc)


@school_required
def api_audit(request):
    if request.method != "GET": return error("Método no permitido.", 405)
    if not allows(request, "school_admin", "directivo"):
        return error("Tu perfil no puede consultar el historial.", 403)
    rows = Audit.objects.filter(school=request.active_school).select_related("actor")[:100]
    return JsonResponse([{"id": a.id, "user": a.actor.name if a.actor else "Sistema", "action": a.action,
                          "detail": a.detail, "created_at": timezone.localtime(a.created_at).strftime("%Y-%m-%d %H:%M:%S")} for a in rows], safe=False)


@school_required
def api_plans(request):
    school = request.active_school
    if request.method == "GET": return JsonResponse([serialize_plan(p) for p in AcademicPlan.objects.filter(school=school)], safe=False)
    if request.method != "POST": return error("Método no permitido.", 405)
    if not allows(request, "school_admin", "directivo"): return error("Sin permiso.", 403)
    try:
        data = json_body(request)
        school_type = data.get("school_type", school.school_type)
        if school_type not in School.Type.values: raise ValueError("Tipo de secundaria inválido.")
        plan = AcademicPlan.objects.create(school=school, name=str(data["name"]).strip(), school_type=school_type,
            jurisdiction=str(data.get("jurisdiction", school.jurisdiction)), orientation=str(data.get("orientation", "")),
            valid_from_year=int(data["valid_from_year"]))
        audit(request, "Plan académico", f"Alta de {plan.name}")
        return JsonResponse(serialize_plan(plan), status=201)
    except (KeyError, ValueError, IntegrityError) as exc: return error(exc)


@school_required
def api_plan_years(request):
    school = request.active_school
    if request.method == "GET":
        qs = PlanYear.objects.filter(school=school).select_related("plan").prefetch_related("subject_entries__subject")
        plan_id = request.GET.get("plan")
        if plan_id: qs = qs.filter(plan_id=plan_id)
        return JsonResponse([serialize_year(y) for y in qs], safe=False)
    if request.method != "POST": return error("Método no permitido.", 405)
    if not allows(request, "school_admin", "directivo"): return error("Sin permiso.", 403)
    try:
        data=json_body(request); plan=AcademicPlan.objects.get(pk=data["plan_id"],school=school)
        subject_rows=data.get("subjects",[])
        subject_ids=[int(item["subject_id"]) for item in subject_rows]
        if len(subject_ids)!=len(set(subject_ids)):raise ValueError("No repitas una materia en el mismo año.")
        subjects={x.id:x for x in Subject.objects.filter(school=school,pk__in=subject_ids)}
        if len(subjects)!=len(subject_ids):raise ValueError("Una materia seleccionada no pertenece a esta escuela.")
        weekly_hours=[float(item.get("weekly_hours",0)) for item in subject_rows]
        if any(hours<0 for hours in weekly_hours):raise ValueError("La carga horaria no puede ser negativa.")
        year=PlanYear.objects.create(school=school,plan=plan,year_label=str(data["year_label"]).strip(),ordinal=int(data["ordinal"]),cycle=str(data.get("cycle","")),orientation=str(data.get("orientation",plan.orientation)))
        for item,hours in zip(subject_rows,weekly_hours):
            PlanSubject.objects.create(school=school,plan_year=year,subject=subjects[int(item["subject_id"])],weekly_hours=hours)
        audit(request,"Año de plan",f"{plan.name} · {year.year_label}")
        return JsonResponse(serialize_year(year),status=201)
    except (KeyError,ValueError,IntegrityError,AcademicPlan.DoesNotExist,Subject.DoesNotExist) as exc:return error(exc)


@school_required
def api_sections(request):
    school=request.active_school
    if request.method=="GET":
        qs=Section.objects.filter(school=school).select_related("plan_year")
        year=request.GET.get("academic_year")
        if year: qs=qs.filter(academic_year=year)
        if active_role(request)=="docente" and not request.user.is_superuser:
            qs=qs.filter(academic_year=timezone.localdate().year,
                offerings__teacher_assignments__membership=request.school_membership).distinct()
        return JsonResponse([serialize_section(s) for s in qs],safe=False)
    if request.method!="POST":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo","secretaria"):return error("Sin permiso.",403)
    try:
        data=json_body(request); py=PlanYear.objects.get(pk=data["plan_year_id"],school=school)
        section=Section.objects.create(school=school,plan_year=py,academic_year=int(data.get("academic_year",timezone.localdate().year)),division=str(data["division"]),shift=data.get("shift",Section.Shift.MORNING),name=str(data.get("name","")))
        return JsonResponse(serialize_section(section),status=201)
    except (KeyError,ValueError,IntegrityError,PlanYear.DoesNotExist) as exc:return error(exc)


@school_required
def api_enrollments(request):
    school=request.active_school
    role=active_role(request)
    if request.method=="GET":
        if role not in WRITE_ROLES|{"secretaria","preceptor","docente","alumno","tutor"} and not request.user.is_superuser:
            return error("Tu perfil no puede consultar inscripciones.",403)
        qs=Enrollment.objects.filter(school=school).select_related("student","section__plan_year")
        if role in {"alumno","tutor","docente"}:
            qs=qs.filter(student__in=student_query(request))
        if role=="docente" and not request.user.is_superuser:
            qs=qs.filter(academic_year=timezone.localdate().year,
                section__offerings__teacher_assignments__membership=request.school_membership).distinct()
        year=request.GET.get("academic_year")
        section=request.GET.get("section")
        if year:qs=qs.filter(academic_year=year)
        if section:qs=qs.filter(section_id=section)
        return JsonResponse([serialize_enrollment(e) for e in qs.order_by("academic_year","student__name")],safe=False)
    if request.method!="POST":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo","secretaria"):return error("Tu perfil no puede administrar inscripciones.",403)
    try:
        data=json_body(request)
        student=Student.objects.get(pk=int(data["student_id"]),school=school)
        section=Section.objects.select_related("plan_year").get(pk=int(data["section_id"]),school=school)
        academic_year=int(data.get("academic_year",section.academic_year))
        if academic_year!=section.academic_year:raise ValueError("El ciclo lectivo debe coincidir con el de la comisión.")
        enrollment,created=Enrollment.objects.update_or_create(school=school,student=student,academic_year=academic_year,
            defaults={"section":section,"state":Enrollment.State.ACTIVE,"started_at":timezone.localdate(),"ended_at":None})
        if created:
            Enrollment.objects.filter(school=school,student=student,state=Enrollment.State.ACTIVE,academic_year__lt=academic_year).exclude(pk=enrollment.pk).update(state=Enrollment.State.PROMOTED,ended_at=timezone.localdate())
        audit(request,"Inscripción",f"{student.name} · {section.name or section.plan_year.year_label} · {academic_year}")
        return JsonResponse(serialize_enrollment(enrollment),status=201 if created else 200)
    except (KeyError,ValueError,TypeError,Student.DoesNotExist,Section.DoesNotExist,IntegrityError) as exc:return error(exc)


@school_required
def api_offerings(request):
    school=request.active_school
    if request.method=="GET":
        qs=Offering.objects.filter(school=school).select_related("section__plan_year","subject").prefetch_related("teacher_assignments__membership__user")
        section=request.GET.get("section")
        if section:qs=qs.filter(section_id=section)
        if active_role(request)=="docente" and not request.user.is_superuser:
            qs=qs.filter(teacher_assignments__membership=request.school_membership,section__academic_year=timezone.localdate().year).distinct()
        return JsonResponse([serialize_offering(o) for o in qs],safe=False)
    if request.method!="POST":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo"):return error("Sin permiso.",403)
    try:
        data=json_body(request); section=Section.objects.get(pk=data["section_id"],school=school); subject=Subject.objects.get(pk=data["subject_id"],school=school)
        membership_ids=[int(mid) for mid in data.get("teacher_membership_ids",[])]
        if len(membership_ids)!=len(set(membership_ids)):raise ValueError("No repitas un docente en la asignación.")
        memberships={m.id:m for m in Membership.objects.filter(pk__in=membership_ids,school=school,is_active=True,role=Membership.Role.TEACHER)}
        if len(memberships)!=len(membership_ids):raise ValueError("Un docente seleccionado no pertenece a esta escuela o no está activo.")
        offering,_=Offering.objects.get_or_create(school=school,section=section,subject=subject)
        for membership in memberships.values():
            TeacherAssignment.objects.get_or_create(school=school,offering=offering,membership=membership)
        audit(request,"Asignación docente",f"{subject.name} · {section.name}")
        return JsonResponse(serialize_offering(offering),status=201)
    except (KeyError,ValueError,IntegrityError,Section.DoesNotExist,Subject.DoesNotExist,Membership.DoesNotExist) as exc:return error(exc)


def parse_import_file(upload):
    if upload.size > 10 * 1024 * 1024:
        raise ValueError("El archivo supera el límite de 10 MB.")
    suffix=upload.name.lower().rsplit(".",1)[-1] if "." in upload.name else ""
    if suffix=="csv":
        text=upload.read().decode("utf-8-sig")
        try:
            dialect=csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect=csv.excel
        return list(csv.reader(io.StringIO(text), dialect=dialect))
    if suffix=="xlsx":
        wb=load_workbook(upload,read_only=True,data_only=True)
        sheet=wb.active
        return [["" if c is None else str(c).strip() for c in row] for row in sheet.iter_rows(values_only=True)]
    raise ValueError("Usá una planilla .csv o .xlsx.")


def normalize_header(value):
    value=unicodedata.normalize("NFKD",str(value or "")).encode("ascii","ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+","_",value).strip("_")


def import_row_error(row, resource, school):
    errors=[]
    if resource=="students":
        if not row.get("name"):errors.append("Falta nombre.")
        if not row.get("course"):errors.append("Falta año/curso.")
        if not row.get("division"):errors.append("Falta división.")
        source=row.get("source_id","")
        if source and Student.objects.filter(school=school,source_id=source).exists():errors.append("El identificador ya está en uso.")
        if row.get("email") and ("@" not in row["email"] or "." not in row["email"].split("@")[-1]):errors.append("Correo inválido.")
        elif row.get("email") and Student.objects.filter(school=school,email__iexact=row["email"]).exists():errors.append("El correo ya está asociado a un alumno de esta escuela.")
    elif resource=="courses":
        if not row.get("name"):errors.append("Falta nombre.")
        if row.get("name") and Course.objects.filter(school=school,name=row["name"]).exists():errors.append("El curso ya existe.")
    elif resource=="subjects":
        if not row.get("name"):errors.append("Falta nombre.")
        if row.get("kind") not in {"subject","workshop"}:errors.append("Tipo debe ser subject o workshop.")
        if row.get("name") and Subject.objects.filter(school=school,name=row["name"],kind=row.get("kind")).exists():errors.append("La materia/taller ya existe.")
    elif resource=="staff":
        if not row.get("name"):errors.append("Falta nombre.")
        if not row.get("email") or "@" not in row["email"]:errors.append("Correo inválido o faltante.")
        if row.get("role") not in {x for x in ROLE_ALIASES if x not in {"alumno","tutor"}}:errors.append("Rol no permitido para importación de personal.")
        if row.get("email") and Membership.objects.filter(school=school,user__email__iexact=row["email"]).exists():errors.append("Ya tiene una cuenta en esta escuela.")
    return errors


@school_required
def api_import_preview(request):
    if request.method!="POST":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo","secretaria"):return error("Sin permiso para importar planillas.",403)
    resource=request.POST.get("resource","students")
    if resource not in {"students","courses","subjects","staff"}:return error("Tipo de planilla no válido.")
    upload=request.FILES.get("file")
    if not upload:return error("Seleccioná un archivo CSV o XLSX.")
    try:
        raw=parse_import_file(upload)
        if len(raw)<1:raise ValueError("El archivo no tiene encabezados.")
        headers=[normalize_header(x) for x in raw[0]]
        aliases={"nombre":"name","alumno":"name","estudiante":"name","correo":"email","curso":"course","ano":"course","año":"course","division":"division","turno":"shift","id":"source_id","identificador":"source_id","codigo":"source_id","código":"source_id","tipo":"kind","materia":"name","taller":"name","rol":"role","perfil":"role","orientacion":"orientation","orientación":"orientation"}
        headers=[aliases.get(x,x) for x in headers]
        allowed={"students":{"source_id","name","email","course","division","shift","orientation"},"courses":{"name"},"subjects":{"name","kind"},"staff":{"name","email","role"}}[resource]
        records=[]; errors=[]; seen=set()
        for line,row_values in enumerate(raw[1:],start=2):
            if not any(str(x or "").strip() for x in row_values):continue
            record={headers[i]:str(row_values[i] or "").strip() for i in range(min(len(headers),len(row_values))) if headers[i] in allowed}
            problems=import_row_error(record,resource,request.active_school)
            if resource=="students":
                keys=[("student_id",record["source_id"].lower())] if record.get("source_id") else []
                if record.get("email"):keys.append(("student_email",record["email"].lower()))
                if not keys and record.get("name") and record.get("course") and record.get("division"):
                    keys.append(("student_class",normalize_header(record["name"])+"|"+normalize_header(record["course"])+"|"+normalize_header(record["division"])))
            elif resource=="staff":keys=[("staff_email",record.get("email","").lower())] if record.get("email") else []
            elif resource=="subjects":keys=[("subject",normalize_header(record.get("name",""))+"|"+record.get("kind",""))] if record.get("name") else []
            else:keys=[(resource,normalize_header(record.get("name","")))] if record.get("name") else []
            if any(key in seen for key in keys):problems.append("Registro duplicado dentro del archivo.")
            seen.update(keys)
            records.append(record)
            if problems:errors.append({"row":line,"messages":problems})
        batch=ImportBatch.objects.create(school=request.active_school,created_by=request.user,resource=resource,rows=records,errors=errors)
        return JsonResponse({"batch_id":str(batch.token),"resource":resource,"rows":len(records),"valid_rows":len(records)-len(errors),"errors":errors,"preview":records[:10]},status=201)
    except (ValueError,UnicodeDecodeError,Exception) as exc:
        return error(exc)


@school_required
def api_import_commit(request, token):
    if request.method!="POST":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo","secretaria"):return error("Sin permiso.",403)
    try:
        with transaction.atomic():
            batch=ImportBatch.objects.select_for_update().get(token=token,school=request.active_school,created_by=request.user,state=ImportBatch.State.PREVIEW)
            errors_by_row={x["row"]:x for x in batch.errors}
            # The batch is created only after row validation; revalidate the accepted rows at commit time.
            created=0; skipped=0; invite_emails=[]
            for index,row in enumerate(batch.rows,start=2):
                if any(e["row"]==index for e in batch.errors):skipped+=1;continue
                current_errors=import_row_error(row,batch.resource,request.active_school)
                if current_errors:skipped+=1;continue
                if batch.resource=="students":
                    create_student(request.active_school,row)
                elif batch.resource=="courses":
                    Course.objects.create(school=request.active_school,name=row["name"])
                elif batch.resource=="subjects":
                    Subject.objects.create(school=request.active_school,name=row["name"],kind=row["kind"])
                elif batch.resource=="staff":
                    user,created_user=User.objects.get_or_create(email=row["email"].lower(),defaults={"name":row["name"]})
                    if not created_user and Membership.objects.filter(school=request.active_school,user=user).exists():skipped+=1;continue
                    if created_user:user.set_unusable_password();user.save(update_fields=("password",))
                    membership=Membership.objects.create(school=request.active_school,user=user,role=row["role"])
                    invite_emails.append(user)
                created+=1
            for user in invite_emails:send_invite(request,user)
            batch.state=ImportBatch.State.COMMITTED
            batch.save(update_fields=("state",))
            audit(request,"Importación",f"{batch.resource}: {created} registros; {skipped} filas omitidas")
        return JsonResponse({"ok":True,"created":created,"skipped":skipped})
    except ImportBatch.DoesNotExist:return error("La vista previa no existe, expiró o ya se confirmó.",404)
    except (IntegrityError,ValueError,Exception) as exc:return error(exc)


@school_required
def api_import_template(request):
    if request.method!="GET":return error("Método no permitido.",405)
    if not allows(request,"school_admin","directivo","secretaria"):return error("Sin permiso.",403)
    resource=request.GET.get("resource","students")
    headers={"students":["source_id","name","email","course","division","shift","orientation"],
        "courses":["name"],"subjects":["name","kind"],"staff":["name","email","role"]}
    if resource not in headers:return error("Tipo de planilla no válido.")
    file_format=request.GET.get("format","csv").lower()
    if file_format=="xlsx":
        workbook=Workbook()
        sheet=workbook.active
        sheet.title=resource[:31]
        sheet.append(headers[resource])
        output=io.BytesIO()
        workbook.save(output)
        response=HttpResponse(output.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"]=f'attachment; filename="nexo-{resource}-template.xlsx"'
        return response
    if file_format!="csv":return error("El formato debe ser CSV o XLSX.")
    response=HttpResponse("\ufeff"+",".join(headers[resource])+"\n",content_type="text/csv; charset=utf-8")
    response["Content-Disposition"]=f'attachment; filename="nexo-{resource}-template.csv"'
    return response


@api_methods("GET")
def api_health(request):
    from django.db import connections
    try:
        connections["default"].ensure_connection()
        return JsonResponse({"ok":True,"database":connections["default"].vendor})
    except Exception:return error("La base de datos no está disponible.",503)


@api_methods("GET")
@school_required
def api_export(request, resource):
    if not allows(request, "school_admin", "directivo", "secretaria"):
        return error("Tu perfil no puede exportar información escolar.", 403)
    school = request.active_school
    if resource == "students":
        headers = ["identificador", "nombre", "correo", "año", "división", "turno", "ciclo"]
        rows = []
        for student in Student.objects.filter(school=school).order_by("name"):
            data = student_data(student, school)
            enrollment = enrollment_for(student, school)
            rows.append([student.source_id, student.name, student.email, data["course"], data["division"], data["shift"], enrollment.academic_year if enrollment else ""])
    elif resource == "attendance":
        headers = ["fecha", "identificador", "alumno", "estado", "materia_taller", "observación"]
        rows = [[a.date.isoformat(), a.student.source_id, a.student.name, a.get_status_display(), a.offering.subject.name if a.offering_id else "Jornada", a.note]
                for a in Attendance.objects.filter(school=school).select_related("student", "offering__subject").order_by("date", "student__name")]
    elif resource == "grades":
        headers = ["año", "período", "identificador", "alumno", "materia_taller", "nota", "observación"]
        rows = [[g.period.year, g.period.name, g.student.source_id, g.student.name, g.offering.subject.name, str(g.value), g.note]
                for g in Grade.objects.filter(school=school).select_related("period", "student", "offering__subject").order_by("period__year", "student__name")]
    elif resource == "loans":
        headers = ["libro", "código", "identificador", "alumno", "retirado", "vencimiento", "devuelto"]
        rows = [[l.book.title, l.book.code, l.student.source_id, l.student.name, l.borrowed_at.isoformat(), l.due_at.isoformat(), l.returned_at.isoformat() if l.returned_at else ""]
                for l in Loan.objects.filter(school=school).select_related("book", "student").order_by("borrowed_at")]
    else:
        return error("Tipo de exportación no disponible.", 404)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([("'" + str(value)) if str(value).startswith(("=", "+", "-", "@", "\t")) else value for value in row])
    response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="nexo-{resource}-{timezone.localdate().isoformat()}.csv"'
    audit(request, "Exportación", f"Descarga de {resource}")
    return response


def password_reset_complete(request):
    return render(request,"core/password_reset_complete.html")


def purge_expired_school_signup_requests():
    return SchoolSignupRequest.objects.filter(expires_at__lte=timezone.now()).delete()[0]


def signup_plan_quotes():
    pricing = PlatformBillingSettings.objects.filter(pk=1).first()
    if not pricing or pricing.onboarding_amount_ars <= 0:
        return []
    first_day = timezone.localdate().replace(day=1)
    plans = []
    for key, label in SubscriptionPlan.choices:
        monthly = monthly_price_for(first_day, key, pricing)
        if monthly > 0:
            plans.append({"id": key, "label": label, "monthly": str(monthly),
                "onboarding": str(pricing.onboarding_amount_ars),
                "initial_total": str(monthly + pricing.onboarding_amount_ars)})
    return plans


@ensure_csrf_cookie
def school_signup_page(request):
    purge_expired_school_signup_requests()
    return render(request, "core/school_signup.html", {"plans": signup_plan_quotes()})


def school_signup_submit(request):
    if request.method != "POST":
        return error("Método no permitido.", 405)
    purge_expired_school_signup_requests()
    try:
        data = json_body(request)
    except ValueError as exc:
        return error(exc)
    # A filled honeypot is silently accepted; it never creates a request or sends mail.
    if str(data.get("website", "")).strip():
        return JsonResponse({"message": "Si los datos son válidos, te enviaremos un correo."}, status=202)
    school_name = str(data.get("school_name", "")).strip()
    school_type = str(data.get("school_type", ""))
    jurisdiction = str(data.get("jurisdiction", "")).strip()
    contact_name = str(data.get("contact_name", "")).strip()
    contact_email = str(data.get("contact_email", "")).strip().lower()
    contact_phone = str(data.get("contact_phone", "")).strip()
    plan = str(data.get("plan", ""))
    if not school_name or len(school_name) > 180 or school_type not in School.Type.values:
        return error("Completá un nombre y tipo de escuela válidos.")
    if len(jurisdiction) > 120 or not contact_name or len(contact_name) > 180:
        return error("Revisá el nombre del director y la jurisdicción.")
    if not contact_phone or len(contact_phone) > 40 or not re.fullmatch(r"[+0-9() .-]{7,40}", contact_phone):
        return error("Ingresá un teléfono válido con característica.")
    try:
        validate_email(contact_email)
    except ValidationError:
        return error("Ingresá un correo electrónico válido.")
    plans = {row["id"]: row for row in signup_plan_quotes()}
    if plan not in plans:
        return error("El plan elegido no está disponible. Actualizá la página y volvé a intentarlo.")

    now = timezone.now()
    ip = request.META.get("REMOTE_ADDR", "")[:64]
    ip_digest = hmac.new(settings.SECRET_KEY.encode(), ip.encode(), hashlib.sha256).hexdigest()
    email_count = SchoolSignupRequest.objects.filter(contact_email=contact_email,
        created_at__gte=now - timedelta(hours=24)).count()
    ip_count = SchoolSignupRequest.objects.filter(ip_digest=ip_digest,
        created_at__gte=now - timedelta(hours=24)).count()
    if email_count >= 3 or ip_count >= 8:
        return error("Alcanzaste el límite de solicitudes por ahora. Intentá nuevamente mañana.", 429)
    if SchoolSignupRequest.objects.filter(contact_email=contact_email, school_name__iexact=school_name,
            expires_at__gt=now, state__in=(SchoolSignupRequest.State.EMAIL_PENDING, SchoolSignupRequest.State.VERIFIED)).exists():
        return JsonResponse({"message": "Ya existe una solicitud reciente. Revisá tu correo para continuar."}, status=202)

    token = secrets.token_urlsafe(32)
    quote = plans[plan]
    signup = SchoolSignupRequest.objects.create(
        school_name=school_name, school_type=school_type, jurisdiction=jurisdiction,
        contact_name=contact_name, contact_email=contact_email, contact_phone=contact_phone,
        plan=plan, monthly_quote_ars=quote["monthly"], onboarding_quote_ars=quote["onboarding"],
        quoted_at=now, quote_expires_at=now + timedelta(days=7),
        verification_token_hash=hashlib.sha256(token.encode()).hexdigest(),
        verification_expires_at=now + timedelta(hours=24), expires_at=now + timedelta(hours=24),
        ip_digest=ip_digest)
    verify_url = request.build_absolute_uri(reverse("school_signup_verify", args=(signup.id, token)))
    subject = "Confirmá el correo para registrar tu escuela en Nexo"
    body = (f"Hola {contact_name},\n\nPara verificar que este correo te pertenece, abrí este enlace y confirmá la solicitud:\n"
        f"{verify_url}\n\nEl enlace vence en 24 horas. Después de verificarlo podrás contactar a Nexo por WhatsApp.\n\n"
        "La solicitud no crea una escuela ni genera cargos hasta que Nexo la revise y apruebe.")
    try:
        sent = send_mail(subject, body, settings.DEFAULT_FROM_EMAIL, [contact_email], fail_silently=False)
        if not sent:
            raise RuntimeError("No se pudo enviar el correo de verificación.")
    except Exception:
        signup.delete()
        return error("No pudimos enviar el correo de verificación. Revisá la configuración de correo e intentá más tarde.", 503)
    return JsonResponse({"message": "Si los datos son válidos, te enviaremos un correo para verificar la dirección."}, status=202)


def school_signup_verify(request, request_id, token):
    if request.method not in {"GET", "POST"}:
        return error("Método no permitido.", 405)
    purge_expired_school_signup_requests()
    signup = SchoolSignupRequest.objects.filter(pk=request_id).first()
    valid = bool(signup and signup.state == SchoolSignupRequest.State.EMAIL_PENDING
        and signup.verification_expires_at > timezone.now()
        and hmac.compare_digest(signup.verification_token_hash, hashlib.sha256(token.encode()).hexdigest()))
    if request.method == "POST" and valid:
        signup.state = SchoolSignupRequest.State.VERIFIED
        signup.verified_at = timezone.now()
        signup.expires_at = timezone.now() + timedelta(days=14)
        signup.verification_token_hash = ""
        signup.save(update_fields=("state", "verified_at", "expires_at", "verification_token_hash"))
        return redirect("school_signup_result", request_id=signup.id)
    return render(request, "core/school_signup_verify.html", {"valid": valid, "signup": signup},
        status=200 if valid else 410)


def school_signup_result(request, request_id):
    purge_expired_school_signup_requests()
    signup = SchoolSignupRequest.objects.filter(pk=request_id, state=SchoolSignupRequest.State.VERIFIED).first()
    whatsapp_url = ""
    if signup:
        raw_number = str(getattr(settings, "NEXO_WHATSAPP_NUMBER", "")).strip()
        digits = re.sub(r"\D", "", raw_number) if re.fullmatch(r"[+0-9() .-]+", raw_number) else ""
        if re.fullmatch(r"\d{8,15}", digits):
            plan_label = dict(SubscriptionPlan.choices).get(signup.plan, signup.plan)
            message = (f"Hola, verifiqué mi solicitud para {signup.school_name}. Código: {signup.id}. "
                f"Plan: {plan_label}. Quisiera coordinar el alta y la capacitación.")
            whatsapp_url = f"https://wa.me/{digits}?{urlencode({'text': message})}"
    return render(request, "core/school_signup_result.html", {"signup": signup, "whatsapp_url": whatsapp_url,
        "initial_total_ars": signup.monthly_quote_ars + signup.onboarding_quote_ars if signup else None})


def platform_school_admin(request):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method == "GET":
        rows = []
        for school in School.objects.all().order_by("name"):
            subscription = getattr(school, "subscription", None)
            if subscription:
                apply_due_subscription_plan(subscription)
            pending = [charge for charge in subscription.charges.all()
                       if charge.state == SubscriptionCharge.State.PENDING] if subscription else []
            rows.append({"id": school.id, "name": school.name, "slug": school.slug,
                "school_type": school.school_type, "jurisdiction": school.jurisdiction,
                "state": school.state, "members": school.memberships.count(),
                "subscription_state": subscription.state if subscription else None,
                "plan": subscription.plan if subscription else None,
                "pending_plan": subscription.pending_plan if subscription else None,
                "plan_change_effective_on": subscription.plan_change_effective_on.isoformat() if subscription and subscription.plan_change_effective_on else None,
                "monthly_amount_ars": str(subscription.monthly_amount_ars) if subscription else None,
                "pending_charges": len(pending),
                "overdue_charges": sum(1 for charge in pending if charge.due_on < timezone.localdate())})
        return JsonResponse(rows, safe=False)
    if request.method == "POST":
        try:
            school = create_platform_school(request, json_body(request))
            subscription = school.subscription
            return JsonResponse({"id": school.id, "name": school.name, "state": school.state,
                "subscription_state": subscription.state, "plan": subscription.plan,
                "initial_charges": list(subscription.charges.values("id", "kind", "plan", "amount_ars", "due_on"))}, status=201)
        except (KeyError, ValueError, IntegrityError) as exc:
            return error(exc)
    return error("Método no permitido.", 405)


def platform_school_signup_requests(request):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "GET":
        return error("Método no permitido.", 405)
    purge_expired_school_signup_requests()
    rows = []
    for signup in SchoolSignupRequest.objects.filter(state=SchoolSignupRequest.State.VERIFIED).order_by("created_at"):
        rows.append({"id": str(signup.id), "school_name": signup.school_name,
            "school_type": signup.school_type, "jurisdiction": signup.jurisdiction,
            "contact_name": signup.contact_name, "contact_email": signup.contact_email,
            "contact_phone": signup.contact_phone, "plan": signup.plan,
            "monthly_quote_ars": str(signup.monthly_quote_ars),
            "onboarding_quote_ars": str(signup.onboarding_quote_ars),
            "initial_total_ars": str(signup.monthly_quote_ars + signup.onboarding_quote_ars),
            "created_at": signup.created_at.isoformat(), "verified_at": signup.verified_at.isoformat() if signup.verified_at else None,
            "quote_expires_at": signup.quote_expires_at.isoformat(),
            "quote_expired": signup.quote_expires_at <= timezone.now()})
    return JsonResponse(rows, safe=False)


@transaction.atomic
def platform_school_signup_approve(request, request_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        signup = SchoolSignupRequest.objects.select_for_update().get(pk=request_id)
        if signup.state != SchoolSignupRequest.State.VERIFIED:
            raise ValueError("La solicitud ya fue atendida o no verificó su correo.")
        now = timezone.now()
        if signup.expires_at <= now:
            signup.delete()
            raise ValueError("La solicitud venció. La persona deberá iniciar una nueva.")
        if signup.quote_expires_at <= now:
            raise ValueError("La cotización venció. Confirmá los valores actuales por WhatsApp y actualizá la cotización.")
        base_slug = slugify(signup.school_name)[:90].strip("-") or "escuela"
        slug = base_slug
        suffix = 2
        while School.objects.filter(slug=slug).exists():
            slug = f"{base_slug[:90-len(str(suffix))-1]}-{suffix}"
            suffix += 1
        school = School.objects.create(name=signup.school_name, slug=slug,
            school_type=signup.school_type, jurisdiction=signup.jurisdiction, state=School.State.ONBOARDING)
        set_rls_school(school)
        GradingScale.objects.create(school=school)
        first_day = timezone.localdate().replace(day=1)
        next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
        subscription = SchoolSubscription.objects.create(school=school, plan=signup.plan,
            monthly_amount_ars=signup.monthly_quote_ars, onboarding_amount_ars=signup.onboarding_quote_ars,
            contact_name=signup.contact_name, contact_email=signup.contact_email)
        SubscriptionCharge.objects.create(subscription=subscription, kind=SubscriptionCharge.Kind.ONBOARDING,
            plan=signup.plan, amount_ars=signup.onboarding_quote_ars, due_on=timezone.localdate())
        SubscriptionCharge.objects.create(subscription=subscription, kind=SubscriptionCharge.Kind.MONTHLY,
            period_start=first_day, period_end=next_month - timedelta(days=1), plan=signup.plan,
            amount_ars=signup.monthly_quote_ars, due_on=timezone.localdate())
        signup.state = SchoolSignupRequest.State.CONVERTED
        signup.converted_school = school
        signup.reviewed_at = now
        signup.reviewed_by = request.user
        signup.save(update_fields=("state", "converted_school", "reviewed_at", "reviewed_by"))
        Audit.objects.create(school=school, actor=request.user, action="Alta aprobada desde solicitud pública",
            detail=f"{school.name} · solicitud {signup.id}")
        return JsonResponse({"id": school.id, "name": school.name, "state": school.state,
            "subscription_state": subscription.state, "plan": subscription.plan,
            "initial_charges": list(subscription.charges.values("id", "kind", "plan", "amount_ars", "due_on"))}, status=201)
    except SchoolSignupRequest.DoesNotExist:
        return error("No se encontró una solicitud vigente.", 404)
    except (ValueError, IntegrityError) as exc:
        if isinstance(exc, IntegrityError):
            transaction.set_rollback(True)
        return error(exc)


@transaction.atomic
def platform_school_signup_reject(request, request_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        signup = SchoolSignupRequest.objects.select_for_update().get(pk=request_id)
        if signup.state != SchoolSignupRequest.State.VERIFIED:
            raise ValueError("La solicitud ya fue atendida o no verificó su correo.")
        data = json_body(request)
        reason = str(data.get("reason", "")).strip()[:300]
        signup.state = SchoolSignupRequest.State.REJECTED
        signup.reviewed_at = timezone.now()
        signup.reviewed_by = request.user
        signup.rejection_reason = reason
        signup.save(update_fields=("state", "reviewed_at", "reviewed_by", "rejection_reason"))
        return JsonResponse({"id": str(signup.id), "state": signup.state})
    except SchoolSignupRequest.DoesNotExist:
        return error("No se encontró la solicitud.", 404)
    except ValueError as exc:
        return error(exc)


@transaction.atomic
def platform_school_signup_requote(request, request_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        signup = SchoolSignupRequest.objects.select_for_update().get(pk=request_id,
            state=SchoolSignupRequest.State.VERIFIED)
        if signup.expires_at <= timezone.now():
            signup.delete()
            raise ValueError("La solicitud venció. La persona deberá iniciar una nueva.")
        if signup.quote_expires_at > timezone.now():
            raise ValueError("La cotización sigue vigente; no se puede cambiar su importe todavía.")
        data = json_body(request)
        if data.get("whatsapp_confirmed") is not True:
            raise ValueError("Confirmá primero los nuevos importes con la persona por WhatsApp.")
        pricing = PlatformBillingSettings.objects.filter(pk=1).first()
        first_day = timezone.localdate().replace(day=1)
        monthly = monthly_price_for(first_day, signup.plan, pricing) if pricing else Decimal("0")
        if monthly <= 0 or not pricing or pricing.onboarding_amount_ars <= 0:
            raise ValueError("No hay precios configurados para actualizar esta cotización.")
        now = timezone.now()
        signup.monthly_quote_ars = monthly
        signup.onboarding_quote_ars = pricing.onboarding_amount_ars
        signup.quoted_at = now
        signup.quote_expires_at = now + timedelta(days=7)
        signup.save(update_fields=("monthly_quote_ars", "onboarding_quote_ars", "quoted_at", "quote_expires_at"))
        return JsonResponse({"monthly_quote_ars": str(monthly), "onboarding_quote_ars": str(pricing.onboarding_amount_ars),
            "quote_expires_at": signup.quote_expires_at.isoformat()})
    except SchoolSignupRequest.DoesNotExist:
        return error("No se encontró una solicitud vigente.", 404)
    except ValueError as exc:
        return error(exc)


@transaction.atomic
def create_platform_school(request, data):
    name = str(data.get("name", "")).strip()
    slug = str(data.get("slug", "")).strip().lower()
    school_type = data.get("school_type", School.Type.COMMON)
    plan = data.get("plan", SubscriptionPlan.BASIC)
    if not name or not slug or school_type not in School.Type.values:
        raise ValueError("Completá nombre, identificador y tipo de escuela válidos.")
    if plan not in SubscriptionPlan.values:
        raise ValueError("Elegí el plan Básico o Pro.")
    if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", slug):
        raise ValueError("El identificador solo puede incluir minúsculas, números y guiones.")
    email = str(data.get("admin_email", "")).strip().lower()
    admin_name = str(data.get("admin_name", "")).strip()
    if not email or not admin_name:
        raise ValueError("Ingresá nombre y correo de la persona responsable de la escuela.")
    try:
        validate_email(email)
    except ValidationError:
        raise ValueError("El correo de administración no es válido.")
    pricing = PlatformBillingSettings.objects.filter(pk=1).first()
    if not pricing or getattr(pricing, f"{plan}_monthly_amount_ars") <= 0 or pricing.onboarding_amount_ars <= 0:
        raise ValueError("Configurá primero el abono del plan elegido y el cargo de alta en Cobros.")
    school = School.objects.create(name=name, slug=slug, school_type=school_type,
        jurisdiction=str(data.get("jurisdiction", "")).strip(), state=School.State.ONBOARDING)
    set_rls_school(school)
    GradingScale.objects.create(school=school)
    first_day = timezone.localdate().replace(day=1)
    next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
    subscription = SchoolSubscription.objects.create(school=school,
        plan=plan, monthly_amount_ars=monthly_price_for(first_day, plan, pricing),
        onboarding_amount_ars=pricing.onboarding_amount_ars,
        contact_name=admin_name, contact_email=email)
    SubscriptionCharge.objects.create(subscription=subscription, kind=SubscriptionCharge.Kind.ONBOARDING,
        plan=plan, amount_ars=subscription.onboarding_amount_ars, due_on=timezone.localdate())
    SubscriptionCharge.objects.create(subscription=subscription, kind=SubscriptionCharge.Kind.MONTHLY,
        period_start=first_day, period_end=next_month - timedelta(days=1),
        plan=plan, amount_ars=subscription.monthly_amount_ars, due_on=timezone.localdate())
    Audit.objects.create(school=school, actor=request.user, action="Alta pendiente de institución", detail=school.name)
    return school


def parse_ars_amount(value, label):
    try:
        raw = Decimal(str(value))
        amount = raw.quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"{label}: ingresá un importe válido en pesos.")
    if not amount.is_finite() or amount <= 0 or amount > Decimal("9999999999.99"):
        raise ValueError(f"{label}: el importe debe ser mayor que cero.")
    if raw != amount:
        raise ValueError(f"{label}: se permiten hasta dos decimales.")
    return amount


def monthly_price_for(period_start, plan=SubscriptionPlan.BASIC, settings_row=None):
    settings_row = settings_row or PlatformBillingSettings.objects.get_or_create(pk=1)[0]
    revision = PlatformPriceChange.objects.filter(plan=plan, effective_on__lte=period_start).order_by("-effective_on").first()
    return revision.monthly_amount_ars if revision else getattr(settings_row, f"{plan}_monthly_amount_ars")


def effective_subscription_plan(subscription, on_date=None):
    on_date = on_date or timezone.localdate()
    if subscription.pending_plan and subscription.plan_change_effective_on <= on_date:
        return subscription.pending_plan
    return subscription.plan


def apply_due_subscription_plan(subscription, on_date=None):
    on_date = on_date or timezone.localdate()
    target_plan = effective_subscription_plan(subscription, on_date)
    if target_plan == subscription.plan:
        return False
    amount = monthly_price_for(subscription.plan_change_effective_on, target_plan)
    if amount <= 0:
        return False
    subscription.plan = target_plan
    subscription.pending_plan = None
    subscription.plan_change_effective_on = None
    subscription.monthly_amount_ars = amount
    subscription.save(update_fields=("plan", "pending_plan", "plan_change_effective_on", "monthly_amount_ars"))
    return True


def parse_optional_ars_amount(value, label):
    if value in (None, ""):
        return Decimal("0.00")
    try:
        raw = Decimal(str(value))
        amount = raw.quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"{label}: ingresá un importe válido en pesos o dejalo sin configurar.")
    if not amount.is_finite() or amount < 0 or amount > Decimal("9999999999.99") or raw != amount:
        raise ValueError(f"{label}: ingresá cero o un importe positivo con hasta dos decimales.")
    return amount


def platform_billing_settings(request):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    settings_row, _ = PlatformBillingSettings.objects.get_or_create(pk=1)
    if request.method == "GET":
        pending_changes = PlatformPriceChange.objects.filter(effective_on__gte=timezone.localdate()).order_by("effective_on")
        return JsonResponse({"basic_monthly_amount_ars": str(settings_row.basic_monthly_amount_ars),
            "pro_monthly_amount_ars": str(settings_row.pro_monthly_amount_ars),
            "onboarding_amount_ars": str(settings_row.onboarding_amount_ars),
            "price_changes": [{"id": row.id, "plan": row.plan, "monthly_amount_ars": str(row.monthly_amount_ars),
                "effective_on": row.effective_on.isoformat(), "notified_at": row.notified_at.isoformat() if row.notified_at else None}
                for row in pending_changes]})
    if request.method != "PATCH":
        return error("Método no permitido.", 405)
    try:
        data = json_body(request)
        basic_amount = parse_ars_amount(data.get("basic_monthly_amount_ars"), "Abono Básico")
        pro_amount = parse_optional_ars_amount(data.get("pro_monthly_amount_ars"), "Abono Pro")
        onboarding_amount = parse_ars_amount(data.get("onboarding_amount_ars"), "Alta y capacitación")
        if (SchoolSubscription.objects.filter(Q(plan=SubscriptionPlan.BASIC) | Q(pending_plan=SubscriptionPlan.BASIC)).exists()
                and basic_amount != settings_row.basic_monthly_amount_ars):
            raise ValueError("El plan Básico ya tiene escuelas; programá su próximo ajuste en lugar de cambiar la tarifa actual.")
        if (SchoolSubscription.objects.filter(Q(plan=SubscriptionPlan.PRO) | Q(pending_plan=SubscriptionPlan.PRO)).exists()
                and pro_amount != settings_row.pro_monthly_amount_ars):
            raise ValueError("El plan Pro ya tiene escuelas; programá su próximo ajuste en lugar de cambiar la tarifa actual.")
        if SchoolSubscription.objects.exists() and onboarding_amount != settings_row.onboarding_amount_ars:
            raise ValueError("El cargo de alta ya tiene escuelas; no se puede modificar retroactivamente.")
        settings_row.basic_monthly_amount_ars = basic_amount
        settings_row.pro_monthly_amount_ars = pro_amount
        settings_row.onboarding_amount_ars = onboarding_amount
        settings_row.save(update_fields=("basic_monthly_amount_ars", "pro_monthly_amount_ars", "onboarding_amount_ars", "updated_at"))
        return JsonResponse({"basic_monthly_amount_ars": str(settings_row.basic_monthly_amount_ars),
            "pro_monthly_amount_ars": str(settings_row.pro_monthly_amount_ars),
            "onboarding_amount_ars": str(settings_row.onboarding_amount_ars)})
    except (ValueError, KeyError) as exc:
        return error(exc)


def platform_schedule_price_change(request):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        data = json_body(request)
        plan = data.get("plan")
        if plan not in SubscriptionPlan.values:
            raise ValueError("Elegí el plan Básico o Pro para ajustar su tarifa.")
        amount = parse_ars_amount(data.get("monthly_amount_ars"), "Abono mensual")
        effective_on = date.fromisoformat(str(data.get("effective_on", "")))
        if effective_on.day != 1 or effective_on < timezone.localdate() + timedelta(days=30):
            raise ValueError("El ajuste debe empezar el primer día de un mes y anunciarse con al menos 30 días.")
        change = PlatformPriceChange.objects.create(plan=plan, monthly_amount_ars=amount,
            effective_on=effective_on, created_by=request.user)
        outgoing = []
        active_subscriptions = SchoolSubscription.objects.filter(state=SchoolSubscription.State.ACTIVE).select_related("school")
        for subscription in active_subscriptions:
            plan_at_effective_date = (subscription.pending_plan
                if subscription.pending_plan and subscription.plan_change_effective_on <= effective_on
                else subscription.plan)
            if plan_at_effective_date != plan:
                continue
            body = (f"Hola {subscription.contact_name},\n\nEl abono mensual de Nexo Escolar para {subscription.school.name} "
                    f"del plan {subscription.get_plan_display()} se actualizará a ${amount:,.2f} ARS desde el {effective_on:%d/%m/%Y}.\n\n"
                    "Este aviso se envía con más de 30 días de anticipación.\n")
            outgoing.append(("Actualización del abono de Nexo Escolar", body,
                settings.DEFAULT_FROM_EMAIL, [subscription.contact_email]))
        sent = send_mass_mail(tuple(outgoing), fail_silently=False) if outgoing else 0
        if sent == len(outgoing):
            change.notified_at = timezone.now()
            change.save(update_fields=("notified_at",))
        return JsonResponse({"id": change.id, "plan": plan, "monthly_amount_ars": str(amount),
            "effective_on": effective_on.isoformat(), "notifications_sent": sent})
    except (ValueError, KeyError, IntegrityError) as exc:
        return error(exc)


def serialize_charge(charge):
    return {"id": charge.id, "school_id": charge.subscription.school_id,
        "school": charge.subscription.school.name, "kind": charge.kind,
        "plan": charge.plan, "plan_label": charge.get_plan_display(),
        "kind_label": charge.get_kind_display(), "period_start": charge.period_start.isoformat() if charge.period_start else None,
        "period_end": charge.period_end.isoformat() if charge.period_end else None,
        "amount_ars": str(charge.amount_ars), "due_on": charge.due_on.isoformat(),
        "state": charge.state, "state_label": charge.get_state_display(),
        "is_overdue": charge.state == SubscriptionCharge.State.PENDING and charge.due_on < timezone.localdate(),
        "invoice_number": charge.invoice_number, "transfer_reference": charge.transfer_reference,
        "paid_at": charge.paid_at.isoformat() if charge.paid_at else None}


@transaction.atomic
def platform_subscription_charges(request):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method == "GET":
        charges = SubscriptionCharge.objects.select_related("subscription__school").order_by("-due_on", "-id")[:500]
        return JsonResponse([serialize_charge(charge) for charge in charges], safe=False)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        data = json_body(request)
        month_value = str(data.get("month", ""))
        try:
            first_day = date.fromisoformat(month_value + "-01")
        except ValueError:
            raise ValueError("Elegí un mes válido para generar los abonos.")
        current_month = timezone.localdate().replace(day=1)
        if first_day != current_month:
            raise ValueError("Los abonos se generan para el mes actual.")
        next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
        last_day = next_month - timedelta(days=1)
        due_on = first_day + timedelta(days=9)
        created = 0
        skipped = 0
        subscriptions = list(SchoolSubscription.objects.filter(state=SchoolSubscription.State.ACTIVE,
            school__state=School.State.ACTIVE).select_related("school").select_for_update())
        work = []
        for subscription in subscriptions:
            if SubscriptionCharge.objects.filter(subscription=subscription,
                    kind=SubscriptionCharge.Kind.MONTHLY, period_start=first_day).exists():
                skipped += 1
                continue
            target_plan = subscription.plan
            if subscription.pending_plan and subscription.plan_change_effective_on <= first_day:
                target_plan = subscription.pending_plan
            amount = monthly_price_for(first_day, target_plan)
            if amount <= 0:
                raise ValueError(f"Configurá la tarifa {dict(SubscriptionPlan.choices)[target_plan]} antes de generar los cargos.")
            work.append((subscription, target_plan, amount))
        for subscription, target_plan, amount in work:
            update_fields = []
            if target_plan != subscription.plan:
                subscription.plan = target_plan
                subscription.pending_plan = None
                subscription.plan_change_effective_on = None
                update_fields.extend(("plan", "pending_plan", "plan_change_effective_on"))
            subscription.monthly_amount_ars = amount
            update_fields.append("monthly_amount_ars")
            subscription.save(update_fields=tuple(update_fields))
            _, was_created = SubscriptionCharge.objects.get_or_create(subscription=subscription,
                kind=SubscriptionCharge.Kind.MONTHLY, period_start=first_day,
                defaults={"period_end": last_day, "plan": subscription.plan,
                    "amount_ars": subscription.monthly_amount_ars, "due_on": due_on})
            created += int(was_created)
            skipped += int(not was_created)
        return JsonResponse({"created": created, "existing": skipped, "month": first_day.strftime("%Y-%m")})
    except (ValueError, KeyError, IntegrityError) as exc:
        return error(exc)


@transaction.atomic
def activate_paid_subscription(request, subscription):
    if subscription.state != SchoolSubscription.State.PENDING:
        return False
    # Platform billing runs without a school selected in the admin session.
    # Set the tenant scope before touching school-protected rows (membership/audit).
    set_rls_school(subscription.school)
    setup = subscription.charges.filter(kind=SubscriptionCharge.Kind.ONBOARDING).first()
    first_month = subscription.charges.filter(kind=SubscriptionCharge.Kind.MONTHLY).order_by("period_start").first()
    if not setup or not first_month or setup.state != SubscriptionCharge.State.PAID or first_month.state != SubscriptionCharge.State.PAID:
        return False
    user, created = User.objects.get_or_create(email=subscription.contact_email,
        defaults={"name": subscription.contact_name, "is_active": True})
    if not user.is_active:
        raise ValueError("La cuenta de administración existe y está desactivada.")
    if created:
        user.set_unusable_password()
        user.save(update_fields=("password",))
    membership, _ = Membership.objects.get_or_create(school=subscription.school, user=user,
        defaults={"role": Membership.Role.SCHOOL_ADMIN})
    if not membership.is_active:
        membership.is_active = True
        membership.save(update_fields=("is_active",))
    subscription.state = SchoolSubscription.State.ACTIVE
    subscription.activated_at = timezone.now()
    subscription.save(update_fields=("state", "activated_at"))
    subscription.school.state = School.State.ACTIVE
    subscription.school.save(update_fields=("state",))
    send_invite(request, user)
    Audit.objects.create(school=subscription.school, actor=request.user,
        action="Suscripción activada", detail=f"Alta y primer abono pagados · {subscription.contact_email}")
    return True


@transaction.atomic
def platform_subscription_charge_detail(request, charge_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "PATCH":
        return error("Método no permitido.", 405)
    try:
        charge = SubscriptionCharge.objects.select_related("subscription__school").get(pk=charge_id)
        data = json_body(request)
        if data.get("state") != SubscriptionCharge.State.PAID:
            raise ValueError("Solo se pueden registrar pagos confirmados.")
        if charge.state == SubscriptionCharge.State.PAID:
            raise ValueError("Este cargo ya figura como pagado.")
        reference = str(data.get("transfer_reference", "")).strip()
        if not reference:
            raise ValueError("Ingresá la referencia de la transferencia confirmada.")
        charge.state = SubscriptionCharge.State.PAID
        charge.transfer_reference = reference[:120]
        charge.invoice_number = str(data.get("invoice_number", "")).strip()[:80]
        charge.paid_at = timezone.now()
        charge.recorded_by = request.user
        charge.save(update_fields=("state", "transfer_reference", "invoice_number", "paid_at", "recorded_by"))
        activated = activate_paid_subscription(request, charge.subscription)
        return JsonResponse({"charge": serialize_charge(charge), "school_activated": activated})
    except (SubscriptionCharge.DoesNotExist, ValueError, KeyError) as exc:
        return error(exc, 404 if isinstance(exc, SubscriptionCharge.DoesNotExist) else 400)


def platform_school_state(request, school_id):
    if request.method!="PATCH":return error("Método no permitido.",405)
    if not request.user.is_superuser:return error("Acceso solo para administración de la plataforma.",403)
    try:
        data=json_body(request);school=School.objects.get(pk=school_id)
        if data.get("state") not in School.State.values:raise ValueError("Estado inválido.")
        subscription = getattr(school, "subscription", None)
        if data["state"] == School.State.TRIAL and school.state == School.State.ONBOARDING:
            raise ValueError("Una alta paga pendiente no puede habilitarse como piloto gratuito.")
        if data["state"] == School.State.ACTIVE and school.state == School.State.ONBOARDING:
            raise ValueError("La escuela se activa al confirmar el pago del alta y del primer abono.")
        if data["state"] in {School.State.ACTIVE, School.State.TRIAL} and subscription and subscription.state != SchoolSubscription.State.ACTIVE:
            raise ValueError("La suscripción de esta escuela no está activa.")
        school.state=data["state"];school.save(update_fields=("state",))
        set_rls_school(school)
        Audit.objects.create(school=school,actor=request.user,action="Estado de institución",detail=school.get_state_display())
        return JsonResponse({"id":school.id,"state":school.state})
    except (School.DoesNotExist,ValueError) as exc:return error(exc,404 if isinstance(exc,School.DoesNotExist) else 400)


def platform_resend_school_invite(request, school_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "POST":
        return error("Método no permitido.", 405)
    try:
        subscription = SchoolSubscription.objects.select_related("school").get(school_id=school_id)
    except SchoolSubscription.DoesNotExist:
        return error("No se encontró la suscripción de esta escuela.", 404)
    if subscription.state != SchoolSubscription.State.ACTIVE or subscription.school.state != School.State.ACTIVE:
        return error("La invitación solo se puede reenviar a una escuela activa.", 400)
    user = User.objects.filter(email__iexact=subscription.contact_email, is_active=True).first()
    if not user or not Membership.objects.filter(school=subscription.school, user=user, is_active=True).exists():
        return error("No se encontró una cuenta activa para la persona responsable de esta escuela.", 400)
    try:
        send_invite(request, user)
    except Exception:
        return error("No se pudo enviar la invitación. Revisá la configuración SMTP y volvé a intentar.", 502)
    Audit.objects.create(school=subscription.school, actor=request.user,
        action="Invitación de acceso reenviada", detail=f"Invitación reenviada a {subscription.contact_email}")
    return JsonResponse({"sent": True})


@transaction.atomic
def platform_school_subscription(request, school_id):
    if not request.user.is_superuser:
        return error("Acceso solo para administración de la plataforma.", 403)
    if request.method != "PATCH":
        return error("Método no permitido.", 405)
    try:
        data = json_body(request)
        subscription = SchoolSubscription.objects.select_related("school").get(school_id=school_id)
        set_rls_school(subscription.school)
        apply_due_subscription_plan(subscription)
        if data.get("state") == SchoolSubscription.State.ACTIVE:
            if subscription.state != SchoolSubscription.State.CANCELED or subscription.school.state != School.State.SUSPENDED:
                raise ValueError("Solo se puede reactivar una suscripción cancelada.")
            access_enabled = bool(subscription.activated_at)
            charge_created = False
            charge = None
            monthly_amount = subscription.monthly_amount_ars
            if access_enabled:
                first_day = timezone.localdate().replace(day=1)
                next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
                monthly_amount = monthly_price_for(first_day, subscription.plan)
                if monthly_amount <= 0:
                    raise ValueError("Configurá el precio mensual de este plan antes de reactivar la suscripción.")
                due_on = max(first_day + timedelta(days=9), timezone.localdate())
                charge, charge_created = SubscriptionCharge.objects.get_or_create(
                    subscription=subscription, kind=SubscriptionCharge.Kind.MONTHLY, period_start=first_day,
                    defaults={"period_end": next_month - timedelta(days=1), "plan": subscription.plan,
                        "amount_ars": monthly_amount, "due_on": due_on})
                if charge.state == SubscriptionCharge.State.VOID:
                    charge.state = SubscriptionCharge.State.PENDING
                    charge.plan = subscription.plan
                    charge.amount_ars = monthly_amount
                    charge.due_on = due_on
                    charge.invoice_number = ""
                    charge.transfer_reference = ""
                    charge.paid_at = None
                    charge.recorded_by = None
                    charge.save(update_fields=("state", "plan", "amount_ars", "due_on", "invoice_number",
                        "transfer_reference", "paid_at", "recorded_by"))
                    charge_created = True
                subscription.state = SchoolSubscription.State.ACTIVE
                subscription.monthly_amount_ars = monthly_amount
                subscription.save(update_fields=("state", "monthly_amount_ars"))
                subscription.school.state = School.State.ACTIVE
                subscription.school.save(update_fields=("state",))
                notification_status = SchoolSubscription.State.ACTIVE
            else:
                subscription.state = SchoolSubscription.State.PENDING
                subscription.save(update_fields=("state",))
                subscription.school.state = School.State.ONBOARDING
                subscription.school.save(update_fields=("state",))
                notification_status = SchoolSubscription.State.PENDING
            Audit.objects.create(school=subscription.school, actor=request.user,
                action="Suscripción reactivada",
                detail="Acceso restablecido desde la consola de plataforma" if access_enabled else "Alta reabierta; pendiente de pagos")
            try:
                send_subscription_status_email(subscription, notification_status, charge)
            except Exception:
                transaction.set_rollback(True)
                return error("No se pudo enviar el aviso de reactivación; la operación no se aplicó. Revisá la configuración SMTP.", 502)
            return JsonResponse({"school_id": school_id, "state": subscription.state,
                "monthly_charge_created": charge_created, "monthly_amount_ars": str(monthly_amount),
                "access_enabled": access_enabled, "notification_sent": True})
        if data.get("state") == SchoolSubscription.State.CANCELED:
            subscription.state = SchoolSubscription.State.CANCELED
            subscription.pending_plan = None
            subscription.plan_change_effective_on = None
            subscription.save(update_fields=("state", "pending_plan", "plan_change_effective_on"))
            subscription.school.state = School.State.SUSPENDED
            subscription.school.save(update_fields=("state",))
            set_rls_school(subscription.school)
            Audit.objects.create(school=subscription.school, actor=request.user,
                action="Suscripción cancelada", detail="Cancelación manual desde la consola de plataforma")
            try:
                send_subscription_status_email(subscription, SchoolSubscription.State.CANCELED)
            except Exception:
                transaction.set_rollback(True)
                return error("No se pudo enviar el aviso de cancelación; la operación no se aplicó. Revisá la configuración SMTP.", 502)
            return JsonResponse({"school_id": school_id, "state": subscription.state,
                "notification_sent": True})

        plan = data.get("plan")
        if plan not in SubscriptionPlan.values:
            raise ValueError("Elegí el plan Básico o Pro, o indicá una cancelación válida.")
        if subscription.state != SchoolSubscription.State.ACTIVE or subscription.school.state != School.State.ACTIVE:
            raise ValueError("Solo se puede programar un cambio en una suscripción activa.")
        if plan == subscription.plan:
            if subscription.pending_plan:
                subscription.pending_plan = None
                subscription.plan_change_effective_on = None
                subscription.save(update_fields=("pending_plan", "plan_change_effective_on"))
                Audit.objects.create(school=subscription.school, actor=request.user,
                    action="Cambio de plan cancelado", detail=f"Se mantiene el plan {subscription.get_plan_display()}")
            return JsonResponse({"school_id": school_id, "plan": subscription.plan,
                "pending_plan": None, "plan_change_effective_on": None})
        today = timezone.localdate().replace(day=1)
        effective_on = (today.replace(day=28) + timedelta(days=4)).replace(day=1)
        if monthly_price_for(effective_on, plan) <= 0:
            raise ValueError(f"Configurá el abono {dict(SubscriptionPlan.choices)[plan]} antes de programar el cambio.")
        subscription.pending_plan = plan
        subscription.plan_change_effective_on = effective_on
        subscription.save(update_fields=("pending_plan", "plan_change_effective_on"))
        Audit.objects.create(school=subscription.school, actor=request.user,
            action="Cambio de plan programado", detail=f"{subscription.get_plan_display()} → {dict(SubscriptionPlan.choices)[plan]} desde {effective_on.isoformat()}")
        return JsonResponse({"school_id": school_id, "plan": subscription.plan,
            "pending_plan": plan, "plan_change_effective_on": effective_on.isoformat()})
    except (SchoolSubscription.DoesNotExist, ValueError, KeyError) as exc:
        return error(exc, 404 if isinstance(exc, SchoolSubscription.DoesNotExist) else 400)
